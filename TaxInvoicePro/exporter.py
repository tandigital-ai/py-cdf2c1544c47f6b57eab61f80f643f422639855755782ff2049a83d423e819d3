# -*- coding: utf-8 -*-
"""
exporter.py - TAX INVOICE PRO v1.2.0
Xuat Excel, xuat ban the hien HTML, chep tep XML goc, mo tep/thu muc.
Tac gia: TDA - 0986509106
"""

import os
import re
import shutil
import subprocess
import sys
from datetime import datetime

import invoice_render
import kho

log = kho.log

TIEU_DE_EXCEL = [
    ("stt", "STT", 6),
    ("ngay_lap", "Ngày lập", 12),
    ("so_hoa_don", "Số hóa đơn", 13),
    ("ky_hieu_mau_so", "Mẫu số", 8),
    ("ky_hieu_hoa_don", "Ký hiệu", 11),
    ("mst_nguoi_ban", "MST người bán", 16),
    ("ten_nguoi_ban", "Tên người bán", 42),
    ("mst_nguoi_mua", "MST người mua", 16),
    ("ten_nguoi_mua", "Tên người mua", 42),
    ("tien_hang", "Tiền hàng", 16),
    ("thue_gtgt", "Thuế GTGT", 14),
    ("tong_tien", "Tổng tiền thanh toán", 18),
    ("loai_hoa_don", "Loại hóa đơn", 20),
    ("trang_thai", "Trạng thái", 12),
    ("ma_cqt", "Mã cơ quan thuế", 26),
    ("ten_tep", "Tệp XML", 34),
]
_COT_SO = {"tien_hang", "thue_gtgt", "tong_tien"}


def _ten_an_toan(s, md="HoaDon"):
    s = re.sub(r"[\\/:*?\"<>|]+", "_", str(s or "")).strip()
    return s or md


def xuat_excel(danh_sach, duong_dan=None, thu_muc=None, tieu_de_bang=None):
    """Xuat danh sach hoa don ra tep .xlsx. Tra ve duong dan tep."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("Chưa cài thư viện openpyxl. Mở Command Prompt, "
                           "bật môi trường ảo rồi chạy: pip install openpyxl")

    if not duong_dan:
        thu_muc = thu_muc or kho.doc_cau_hinh().get("thu_muc_xuat") or os.getcwd()
        os.makedirs(thu_muc, exist_ok=True)
        duong_dan = os.path.join(
            thu_muc, "BangKeHoaDon_%s.xlsx"
            % datetime.now().strftime("%Y%m%d_%H%M%S"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Bang ke hoa don"
    dam = Font(bold=True, color="FFFFFF", size=10)
    nen = PatternFill("solid", fgColor="1565C0")
    vien = Border(*[Side(style="thin", color="B7C4D4")] * 4)
    giua = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws["A1"] = "%s — BẢNG KÊ HÓA ĐƠN ĐIỆN TỬ" % kho.APP_NAME
    ws["A1"].font = Font(bold=True, size=14, color="0D47A1")
    ws.merge_cells(start_row=1, start_column=1, end_row=1,
                   end_column=len(TIEU_DE_EXCEL))
    ws["A2"] = (tieu_de_bang or "Kết xuất lúc %s — Tổng số: %d hóa đơn"
                % (datetime.now().strftime("%d/%m/%Y %H:%M"), len(danh_sach)))
    ws["A2"].font = Font(italic=True, size=10, color="5B6B7C")
    ws.merge_cells(start_row=2, start_column=1, end_row=2,
                   end_column=len(TIEU_DE_EXCEL))

    for i, (_k, nhan, rong) in enumerate(TIEU_DE_EXCEL, 1):
        o = ws.cell(row=4, column=i, value=nhan)
        o.font = dam
        o.fill = nen
        o.border = vien
        o.alignment = giua
        ws.column_dimensions[get_column_letter(i)].width = rong

    for r, inv in enumerate(danh_sach, start=5):
        for i, (k, _n, _w) in enumerate(TIEU_DE_EXCEL, 1):
            if k == "stt":
                gt = r - 4
            elif k == "ngay_lap":
                gt = kho.fmt_ngay(inv.get("ngay_lap"))
            elif k in _COT_SO:
                gt = kho.to_int(inv.get(k))
            else:
                gt = inv.get(k) or ""
            o = ws.cell(row=r, column=i, value=gt)
            o.border = vien
            if k in _COT_SO:
                o.number_format = "#,##0"
                o.alignment = Alignment(horizontal="right")
            elif k in ("ten_nguoi_ban", "ten_nguoi_mua", "ten_tep"):
                o.alignment = Alignment(horizontal="left", wrap_text=False)
            else:
                o.alignment = Alignment(horizontal="center")

    cuoi = 4 + len(danh_sach) + 1
    ws.cell(row=cuoi, column=9, value="CỘNG").font = Font(bold=True)
    for k in ("tien_hang", "thue_gtgt", "tong_tien"):
        cot = [i for i, (kk, _n, _w) in enumerate(TIEU_DE_EXCEL, 1) if kk == k][0]
        o = ws.cell(row=cuoi, column=cot,
                    value=sum(kho.to_int(x.get(k)) for x in danh_sach))
        o.font = Font(bold=True)
        o.number_format = "#,##0"
        o.border = vien
    ws.cell(row=cuoi + 2, column=1,
            value="Tác giả: %s — Điện thoại: %s" % (kho.TAC_GIA, kho.DIEN_THOAI)
            ).font = Font(italic=True, size=9, color="5B6B7C")

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A4:%s%d" % (
        __import__("openpyxl").utils.get_column_letter(len(TIEU_DE_EXCEL)),
        4 + len(danh_sach))
    wb.save(duong_dan)
    log.info("Đã xuất Excel %d dòng: %s", len(danh_sach), duong_dan)
    return duong_dan


# ----------------------------------------------------------------------
def chep_xml(inv, thu_muc):
    """Chep tep XML goc cua mot hoa don sang thu muc nguoi dung chon."""
    goc = inv.get("duong_dan_xml") or ""
    if not goc or not os.path.isfile(goc):
        return ""
    os.makedirs(thu_muc, exist_ok=True)
    ten = "HoaDon_%s_%s.xml" % (
        _ten_an_toan(inv.get("ky_hieu_day_du") or inv.get("ky_hieu_hoa_don")),
        _ten_an_toan(inv.get("so_hoa_don"), "0"))
    dich = os.path.join(thu_muc, ten)
    t, d = os.path.splitext(dich)
    i = 1
    while os.path.exists(dich):
        dich = "%s_%d%s" % (t, i, d)
        i += 1
    shutil.copy2(goc, dich)
    return dich


def chep_nhieu_xml(danh_sach, thu_muc):
    ok, thieu = [], []
    for inv in danh_sach:
        try:
            d = chep_xml(inv, thu_muc)
        except (OSError, shutil.Error) as e:
            log.warning("Không chép XML: %s", e)
            d = ""
        (ok if d else thieu).append(inv)
    return ok, thieu


def xuat_html(inv, thu_muc):
    return invoice_render.luu_html(inv, thu_muc)


def xuat_nhieu_html(danh_sach, thu_muc):
    os.makedirs(thu_muc, exist_ok=True)
    return invoice_render.luu_nhieu_html(danh_sach, thu_muc)


def mo_duong_dan(duong_dan):
    """Mo tep hoac thu muc bang chuong trinh mac dinh cua he dieu hanh."""
    if not duong_dan or not os.path.exists(duong_dan):
        return False
    try:
        if os.name == "nt":
            os.startfile(duong_dan)              # noqa: S606
        elif sys.platform == "darwin":
            subprocess.Popen(["open", duong_dan])
        else:
            subprocess.Popen(["xdg-open", duong_dan])
        return True
    except OSError as e:
        log.warning("Không mở được %s: %s", duong_dan, e)
        return False
