# -*- coding: utf-8 -*-
"""
cong_cu.py - TAX INVOICE PRO v1.3.0
Cua so "Doi chieu & Bao cao" - chay doc lap, khong sua app.py v1.2.0.
Chay bang: python cong_cu.py  (hoac ChayCongCu.bat)
Tac gia: TDA - 0986509106
"""
import os
import datetime
import threading
import traceback
import webbrowser
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

try:
    import va_ttk  # va loi command=None tren Python 3.12+
except Exception:
    pass

import customtkinter as ctk

import doi_chieu as dc
import bao_cao as bc

XANH, XANH_DAM, NEN = "#1565C0", "#0D47A1", "#F4F7FB"
LUC, CAM, DO, XAM = "#2E7D32", "#EF6C00", "#C62828", "#607D8B"
THU_MUC = os.path.dirname(os.path.abspath(__file__))


def hom_nay():
    return datetime.date.today().strftime("%d/%m/%Y")


class CongCu(ctk.CTk):
    def __init__(self):
        super().__init__()
        ctk.set_appearance_mode("light")
        self.title("TAX INVOICE PRO v1.3.0 — Đối chiếu & Báo cáo")
        self.geometry("1180x780")
        self.minsize(1000, 640)
        self.configure(fg_color=NEN)

        self.ds = []
        self.canh_bao = []
        self.kq_doi_chieu = None

        self._style()
        self._dau_trang()
        self._than()
        self._chan_trang()
        self.nap_kho()

    # ------------------------------------------------------- khung
    def _style(self):
        s = ttk.Style()
        try:
            s.theme_use("clam")
        except Exception:
            pass
        s.configure("CC.Treeview", rowheight=26, font=("Segoe UI", 9),
                    fieldbackground="white", background="white")
        s.configure("CC.Treeview.Heading", font=("Segoe UI", 9, "bold"),
                    background=XANH, foreground="white", relief="flat")
        s.map("CC.Treeview.Heading", background=[("active", XANH_DAM)])

    def _dau_trang(self):
        h = ctk.CTkFrame(self, height=52, corner_radius=0, fg_color=XANH)
        h.pack(fill="x")
        h.pack_propagate(False)
        ctk.CTkLabel(h, text="◆  ĐỐI CHIẾU & BÁO CÁO",
                     font=ctk.CTkFont("Segoe UI", 17, "bold"),
                     text_color="white").pack(side="left", padx=18)
        ctk.CTkLabel(h, text="Giai đoạn 4 & 6 — kiểm tra nghiệp vụ, bảng kê, "
                             "gói lưu trữ",
                     font=ctk.CTkFont("Segoe UI", 11),
                     text_color="#CFE3FA").pack(side="left")
        self.lb_kho = ctk.CTkLabel(h, text="", font=ctk.CTkFont("Segoe UI", 11),
                                   text_color="white")
        self.lb_kho.pack(side="right", padx=18)

    def _chan_trang(self):
        f = ctk.CTkFrame(self, height=30, corner_radius=0, fg_color="#E8EEF6")
        f.pack(fill="x", side="bottom")
        f.pack_propagate(False)
        self.lb_tt = ctk.CTkLabel(f, text="Sẵn sàng.",
                                  font=ctk.CTkFont("Segoe UI", 10),
                                  text_color="#44546A")
        self.lb_tt.pack(side="left", padx=14)
        ctk.CTkLabel(f, text="Tác giả: TDA   •   Điện thoại: 0986509106   •   "
                             "Phiên bản 1.3.0",
                     font=ctk.CTkFont("Segoe UI", 10, "bold"),
                     text_color=XANH_DAM).pack(side="right", padx=14)

    def _than(self):
        self.tab = ctk.CTkTabview(self, fg_color="white",
                                  segmented_button_selected_color=XANH,
                                  segmented_button_selected_hover_color=XANH_DAM)
        self.tab.pack(fill="both", expand=True, padx=12, pady=10)
        for t in ("Kiểm tra dữ liệu", "Đối chiếu bảng kê Cổng Thuế",
                  "Báo cáo & Gói lưu trữ", "So sánh hai hóa đơn"):
            self.tab.add(t)
        self._tab_kiem_tra(self.tab.tab("Kiểm tra dữ liệu"))
        self._tab_doi_chieu(self.tab.tab("Đối chiếu bảng kê Cổng Thuế"))
        self._tab_bao_cao(self.tab.tab("Báo cáo & Gói lưu trữ"))
        self._tab_so_sanh(self.tab.tab("So sánh hai hóa đơn"))

    def _bang(self, cha, cot, rong, cao=14):
        khung = ctk.CTkFrame(cha, fg_color="white")
        khung.pack(fill="both", expand=True, padx=4, pady=(6, 4))
        tv = ttk.Treeview(khung, columns=[c[0] for c in cot], show="headings",
                          height=cao, style="CC.Treeview")
        for (ma, nhan), w in zip(cot, rong):
            tv.heading(ma, text=nhan)
            tv.column(ma, width=w, anchor="w",
                      stretch=(w >= 200))
        dọc = ttk.Scrollbar(khung, orient="vertical", command=tv.yview)
        ngang = ttk.Scrollbar(khung, orient="horizontal", command=tv.xview)
        tv.configure(yscrollcommand=dọc.set, xscrollcommand=ngang.set)
        tv.grid(row=0, column=0, sticky="nsew")
        dọc.grid(row=0, column=1, sticky="ns")
        ngang.grid(row=1, column=0, sticky="ew")
        khung.grid_rowconfigure(0, weight=1)
        khung.grid_columnconfigure(0, weight=1)
        tv.tag_configure("nang", background="#FFEBEE")
        tv.tag_configure("vua", background="#FFF8E1")
        tv.tag_configure("nhe", background="#F1F8E9")
        return tv

    def _nut(self, cha, text, lenh, mau=XANH, w=150):
        return ctk.CTkButton(cha, text=text, command=lenh, width=w, height=32,
                             fg_color=mau, hover_color=XANH_DAM,
                             font=ctk.CTkFont("Segoe UI", 12, "bold"))

    def bao(self, t):
        self.lb_tt.configure(text=t)
        self.update_idletasks()

    # ------------------------------------------------- tab 1: kiem tra
    def _tab_kiem_tra(self, cha):
        thanh = ctk.CTkFrame(cha, fg_color="transparent")
        thanh.pack(fill="x", padx=4, pady=(8, 0))
        self._nut(thanh, "Chạy kiểm tra", self.chay_kiem_tra, LUC).pack(side="left")
        self._nut(thanh, "Nạp lại kho", self.nap_kho, XAM, 120).pack(side="left", padx=8)
        self._nut(thanh, "Xuất cảnh báo ra Excel", self.xuat_canh_bao,
                  CAM, 200).pack(side="left")
        self.lb_cb = ctk.CTkLabel(thanh, text="", font=ctk.CTkFont("Segoe UI", 12,
                                                                   "bold"))
        self.lb_cb.pack(side="right", padx=6)
        self.tv_cb = self._bang(
            cha, [("muc", "Mức"), ("loai", "Loại"), ("hd", "Hóa đơn"),
                  ("mo_ta", "Nội dung cần xử lý")],
            [70, 150, 240, 640], 20)

    def chay_kiem_tra(self):
        self.bao("Đang kiểm tra dữ liệu…")
        try:
            self.canh_bao = dc.tong_hop_canh_bao(self.ds)
        except Exception as e:
            messagebox.showerror("Lỗi", "Không kiểm tra được:\n%s" % e)
            return
        self.tv_cb.delete(*self.tv_cb.get_children())
        dem = {"Nặng": 0, "Vừa": 0, "Nhẹ": 0}
        tag = {"Nặng": "nang", "Vừa": "vua", "Nhẹ": "nhe"}
        for c in self.canh_bao:
            dem[c["muc"]] = dem.get(c["muc"], 0) + 1
            self.tv_cb.insert("", "end",
                              values=(c["muc"], c["loai"], c["hoa_don"],
                                      c["mo_ta"]),
                              tags=(tag.get(c["muc"], ""),))
        if not self.canh_bao:
            self.tv_cb.insert("", "end",
                              values=("", "", "", "Không có cảnh báo — "
                                                  "dữ liệu trong kho sạch."))
        self.lb_cb.configure(
            text="Nặng: %d   Vừa: %d   Nhẹ: %d" % (dem.get("Nặng", 0),
                                                   dem.get("Vừa", 0),
                                                   dem.get("Nhẹ", 0)),
            text_color=DO if dem.get("Nặng") else XANH_DAM)
        self.bao("Đã kiểm tra %d hóa đơn, phát hiện %d cảnh báo."
                 % (len(self.ds), len(self.canh_bao)))

    def xuat_canh_bao(self):
        if not self.canh_bao:
            self.chay_kiem_tra()
        d = filedialog.asksaveasfilename(
            title="Lưu danh sách cảnh báo", defaultextension=".xlsx",
            initialfile="CanhBao_%s.xlsx"
                        % datetime.datetime.now().strftime("%Y%m%d_%H%M"),
            filetypes=[("Excel", "*.xlsx")])
        if not d:
            return
        try:
            bc.xuat_bao_cao(self.ds, d, self.canh_bao)
            self.bao("Đã lưu: %s" % d)
            if messagebox.askyesno("Xong", "Đã lưu tệp:\n%s\n\nMở ngay?" % d):
                os.startfile(d)
        except Exception as e:
            messagebox.showerror("Lỗi", str(e))

    # ------------------------------------------- tab 2: doi chieu Excel
    def _tab_doi_chieu(self, cha):
        thanh = ctk.CTkFrame(cha, fg_color="transparent")
        thanh.pack(fill="x", padx=4, pady=(8, 0))
        ctk.CTkLabel(thanh, text="Bảng kê xuất từ Cổng Thuế:",
                     font=ctk.CTkFont("Segoe UI", 12)).pack(side="left")
        self.o_excel = ctk.CTkEntry(thanh, width=520, height=32,
                                    placeholder_text="Chọn tệp .xlsx hoặc .csv "
                                                     "tải từ hoadondientu.gdt.gov.vn")
        self.o_excel.pack(side="left", padx=8)
        self._nut(thanh, "Chọn tệp…", self.chon_excel, XAM, 110).pack(side="left")
        self._nut(thanh, "ĐỐI CHIẾU", self.chay_doi_chieu, LUC,
                  130).pack(side="left", padx=8)
        self._nut(thanh, "Xuất kết quả", self.xuat_doi_chieu, CAM,
                  130).pack(side="left")

        self.lb_dc = ctk.CTkLabel(cha, text="Chưa đối chiếu.",
                                  font=ctk.CTkFont("Segoe UI", 12, "bold"),
                                  text_color=XANH_DAM)
        self.lb_dc.pack(anchor="w", padx=10, pady=(8, 0))
        self.tv_dc = self._bang(
            cha, [("nhom", "Nhóm"), ("ngay", "Ngày lập"), ("kh", "Ký hiệu"),
                  ("so", "Số HĐ"), ("mst", "MST bán"), ("ten", "Tên người bán"),
                  ("tien", "Tổng tiền"), ("ghi", "Ghi chú")],
            [180, 90, 90, 90, 110, 300, 110, 240], 18)

    def chon_excel(self):
        d = filedialog.askopenfilename(
            title="Chọn bảng kê hóa đơn xuất từ Cổng Thuế",
            filetypes=[("Excel / CSV", "*.xlsx *.xls *.csv"), ("Tất cả", "*.*")])
        if d:
            self.o_excel.delete(0, "end")
            self.o_excel.insert(0, d)

    def chay_doi_chieu(self):
        d = self.o_excel.get().strip().strip('"')
        if not d or not os.path.exists(d):
            messagebox.showwarning("Thiếu tệp", "Hãy chọn tệp bảng kê đã xuất "
                                                "từ Cổng Thuế trước.")
            return
        self.bao("Đang đọc bảng kê…")
        try:
            ds_cong, tt = dc.doc_bang_ke(d)
        except Exception as e:
            messagebox.showerror("Không đọc được bảng kê", str(e))
            self.bao("Lỗi đọc bảng kê.")
            return
        kq = dc.doi_chieu_bang_ke(self.ds, ds_cong)
        self.kq_doi_chieu = kq
        self.tv_dc.delete(*self.tv_dc.get_children())

        for h in kq["chi_co_cong"]:
            self.tv_dc.insert("", "end", tags=("nang",), values=(
                "Cổng có — máy CHƯA có", dc.ngay_vn(h["ngay_lap"]),
                h["ky_hieu_hoa_don"], h["so_hoa_don"], h["mst_nguoi_ban"],
                h["ten_nguoi_ban"], dc.fmt_tien(h["tong_tien"]),
                "Cần tải XML hóa đơn này về máy"))
        for h in kq["chi_co_may"]:
            self.tv_dc.insert("", "end", tags=("vua",), values=(
                "Máy có — cổng KHÔNG có", dc.ngay_vn(h["ngay_lap"]),
                h["ky_hieu_hoa_don"], h["so_hoa_don"], h["mst_nguoi_ban"],
                h["ten_nguoi_ban"], dc.fmt_tien(h["tong_tien"]),
                "Kiểm tra kỳ lọc trên cổng hoặc hóa đơn đã bị hủy"))
        for l in kq["lech_tien"]:
            a, b = l["may"], l["cong"]
            self.tv_dc.insert("", "end", tags=("vua",), values=(
                "Lệch tổng tiền", dc.ngay_vn(a["ngay_lap"]),
                a["ky_hieu_hoa_don"], a["so_hoa_don"], a["mst_nguoi_ban"],
                a["ten_nguoi_ban"], dc.fmt_tien(a["tong_tien"]),
                "Cổng ghi %s đ — lệch %s đ"
                % (dc.fmt_tien(b["tong_tien"]),
                   dc.fmt_tien(b["tong_tien"] - a["tong_tien"]))))
        if not self.tv_dc.get_children():
            self.tv_dc.insert("", "end", values=(
                "Khớp hoàn toàn", "", "", "", "", "", "",
                "Không có chênh lệch nào"))
        self.lb_dc.configure(
            text="Bảng kê: %d dòng (tiêu đề ở dòng %d)  •  Khớp: %d  •  "
                 "Cổng có máy chưa: %d  •  Máy có cổng chưa: %d  •  "
                 "Lệch tiền: %d"
                 % (tt["so_dong"], tt["dong_tieu_de"], kq["khop"],
                    len(kq["chi_co_cong"]), len(kq["chi_co_may"]),
                    len(kq["lech_tien"])))
        self.bao("Đối chiếu xong.")

    def xuat_doi_chieu(self):
        if not self.kq_doi_chieu:
            messagebox.showinfo("Chưa có dữ liệu", "Hãy chạy đối chiếu trước.")
            return
        d = filedialog.asksaveasfilename(
            title="Lưu kết quả đối chiếu", defaultextension=".xlsx",
            initialfile="DoiChieu_CongThue_%s.xlsx"
                        % datetime.datetime.now().strftime("%Y%m%d_%H%M"),
            filetypes=[("Excel", "*.xlsx")])
        if not d:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            wb = Workbook()
            ws = wb.active
            ws.title = "Đối chiếu"
            ws.append(["Nhóm", "Ngày lập", "Ký hiệu", "Số HĐ", "MST bán",
                       "Tên người bán", "Tổng tiền", "Ghi chú"])
            for i, w in enumerate([24, 12, 10, 12, 15, 40, 15, 42], 1):
                ws.column_dimensions[chr(64 + i)].width = w
                o = ws.cell(row=1, column=i)
                o.font = Font(bold=True, color="FFFFFF")
                o.fill = PatternFill("solid", fgColor="1565C0")
            for iid in self.tv_dc.get_children():
                ws.append(list(self.tv_dc.item(iid)["values"]))
            ws.freeze_panes = "A2"
            wb.save(d)
            self.bao("Đã lưu: %s" % d)
            if messagebox.askyesno("Xong", "Đã lưu:\n%s\n\nMở ngay?" % d):
                os.startfile(d)
        except Exception as e:
            messagebox.showerror("Lỗi", str(e))

    # --------------------------------------------- tab 3: bao cao
    def _tab_bao_cao(self, cha):
        k = ctk.CTkFrame(cha, fg_color="white")
        k.pack(fill="x", padx=10, pady=12)
        ctk.CTkLabel(k, text="Kỳ báo cáo",
                     font=ctk.CTkFont("Segoe UI", 13, "bold"),
                     text_color=XANH_DAM).grid(row=0, column=0, columnspan=4,
                                               sticky="w", padx=12, pady=(10, 4))
        ctk.CTkLabel(k, text="Từ ngày (dd/mm/yyyy)").grid(row=1, column=0,
                                                          padx=12, pady=8,
                                                          sticky="w")
        self.o_tu = ctk.CTkEntry(k, width=160, height=32)
        self.o_tu.grid(row=1, column=1, pady=8, sticky="w")
        ctk.CTkLabel(k, text="Đến ngày (dd/mm/yyyy)").grid(row=1, column=2,
                                                           padx=18, sticky="w")
        self.o_den = ctk.CTkEntry(k, width=160, height=32)
        self.o_den.grid(row=1, column=3, pady=8, sticky="w")
        self.o_den.insert(0, hom_nay())
        ctk.CTkLabel(k, text="Để trống cả hai ô nếu muốn lấy toàn bộ kho.",
                     text_color="#667",
                     font=ctk.CTkFont("Segoe UI", 11)).grid(
            row=2, column=0, columnspan=4, sticky="w", padx=12, pady=(0, 10))

        t = ctk.CTkFrame(cha, fg_color="transparent")
        t.pack(fill="x", padx=10)
        self._nut(t, "Xem trước số liệu", self.xem_truoc, XAM, 170).pack(side="left")
        self._nut(t, "XUẤT BẢNG KÊ EXCEL", self.xuat_bang_ke, LUC,
                  210).pack(side="left", padx=8)
        self._nut(t, "XUẤT GÓI LƯU TRỮ", self.xuat_goi, CAM,
                  200).pack(side="left")

        self.txt_bc = tk.Text(cha, height=18, font=("Consolas", 10),
                              bg="white", fg="#1b2733", relief="solid", bd=1,
                              wrap="word")
        self.txt_bc.pack(fill="both", expand=True, padx=10, pady=12)
        self.txt_bc.insert("1.0",
                           "Chọn kỳ rồi bấm 'Xem trước số liệu'.\n\n"
                           "• Bảng kê Excel gồm 6 sheet: Thông tin, Danh sách, "
                           "Dòng hàng, Theo nhà cung cấp, Theo thuế suất, Cảnh báo.\n"
                           "• Gói lưu trữ tạo một thư mục chứa XML gốc, bản thể "
                           "hiện HTML và tệp danh mục Excel — dùng để lưu trữ "
                           "theo quy định.")

    def _ds_ky(self):
        tu = dc.ngay_chuan(self.o_tu.get().strip())
        den = dc.ngay_chuan(self.o_den.get().strip())
        return [h for h in self.ds
                if (not tu or (h["ngay_lap"] and h["ngay_lap"] >= tu))
                and (not den or (h["ngay_lap"] and h["ngay_lap"] <= den))], tu, den

    def xem_truoc(self):
        ds, tu, den = self._ds_ky()
        ncc = {}
        for h in ds:
            ncc[h["mst_nguoi_ban"]] = ncc.get(h["mst_nguoi_ban"], 0) + 1
        cb = dc.tong_hop_canh_bao(ds)
        self.txt_bc.delete("1.0", "end")
        self.txt_bc.insert("1.0",
                           "KỲ: %s  →  %s\n"
                           "Số hóa đơn       : %d\n"
                           "Số nhà cung cấp  : %d\n"
                           "Tổng tiền hàng   : %s đ\n"
                           "Tổng thuế GTGT   : %s đ\n"
                           "Tổng thanh toán  : %s đ\n"
                           "Số cảnh báo      : %d (nặng: %d)\n\n"
                           "Các nhà cung cấp lớn nhất kỳ này:\n"
                           % (dc.ngay_vn(tu) or "đầu kỳ",
                              dc.ngay_vn(den) or "cuối kỳ", len(ds), len(ncc),
                              dc.fmt_tien(sum(h["tien_hang"] for h in ds)),
                              dc.fmt_tien(sum(h["thue_gtgt"] for h in ds)),
                              dc.fmt_tien(sum(h["tong_tien"] for h in ds)),
                              len(cb),
                              len([c for c in cb if c["muc"] == "Nặng"])))
        tong = {}
        for h in ds:
            g = tong.setdefault(h["ten_nguoi_ban"] or h["mst_nguoi_ban"], [0, 0])
            g[0] += 1
            g[1] += h["tong_tien"]
        for ten, g in sorted(tong.items(), key=lambda x: -x[1][1])[:15]:
            self.txt_bc.insert("end", "  %-46s %3d HĐ   %>14s đ\n"
                                      .replace("%>14s", "%14s")
                               % (ten[:46], g[0], dc.fmt_tien(g[1])))
        self.bao("Xem trước %d hóa đơn." % len(ds))

    def xuat_bang_ke(self):
        ds, tu, den = self._ds_ky()
        if not ds:
            messagebox.showinfo("Không có dữ liệu",
                                "Kỳ đã chọn không có hóa đơn nào.")
            return
        d = filedialog.asksaveasfilename(
            title="Lưu bảng kê hóa đơn", defaultextension=".xlsx",
            initialfile="BangKeHoaDon_%s.xlsx"
                        % datetime.datetime.now().strftime("%Y%m%d_%H%M"),
            filetypes=[("Excel", "*.xlsx")])
        if not d:
            return
        self.bao("Đang lập bảng kê…")
        try:
            bc.xuat_bao_cao(ds, d, dc.tong_hop_canh_bao(ds), tu, den)
        except Exception as e:
            messagebox.showerror("Lỗi", "%s\n\n%s" % (e, traceback.format_exc()))
            return
        self.bao("Đã lưu bảng kê: %s" % d)
        if messagebox.askyesno("Xong", "Đã lập bảng kê %d hóa đơn:\n%s\n\nMở ngay?"
                                       % (len(ds), d)):
            os.startfile(d)

    def xuat_goi(self):
        ds, tu, den = self._ds_ky()
        if not ds:
            messagebox.showinfo("Không có dữ liệu",
                                "Kỳ đã chọn không có hóa đơn nào.")
            return
        t = filedialog.askdirectory(title="Chọn nơi đặt gói lưu trữ")
        if not t:
            return
        self.bao("Đang tạo gói lưu trữ…")

        def chay():
            try:
                kq = bc.xuat_goi_luu_tru(
                    ds, t, tu, den,
                    bao_tien_do=lambda i, n: self.bao(
                        "Đang xử lý %d/%d hóa đơn…" % (i, n)))
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Lỗi", str(e)))
                return
            def xong():
                self.bao("Đã tạo gói: %s" % kq["thu_muc"])
                if messagebox.askyesno(
                        "Hoàn tất",
                        "Đã tạo gói lưu trữ:\n%s\n\n"
                        "Hóa đơn: %d\nChép được XML gốc: %d\n"
                        "Thiếu tệp XML: %d\nBản thể hiện HTML: %d\n\n"
                        "Mở thư mục ngay?"
                        % (kq["thu_muc"], len(ds), kq["xml"], kq["thieu_xml"],
                           kq["html"])):
                    os.startfile(kq["thu_muc"])
            self.after(0, xong)

        threading.Thread(target=chay, daemon=True).start()

    # ------------------------------------------- tab 4: so sanh 2 HD
    def _tab_so_sanh(self, cha):
        t = ctk.CTkFrame(cha, fg_color="transparent")
        t.pack(fill="x", padx=10, pady=(10, 0))
        ctk.CTkLabel(t, text="Lọc nhanh (số hóa đơn, ký hiệu hoặc tên "
                             "người bán):").pack(side="left")
        self.o_loc = ctk.CTkEntry(t, width=330, height=32)
        self.o_loc.pack(side="left", padx=8)
        self.o_loc.bind("<Return>", lambda e: self.loc_so_sanh())
        self._nut(t, "Lọc", self.loc_so_sanh, XAM, 90).pack(side="left")
        self._nut(t, "SO SÁNH HAI DÒNG ĐANG CHỌN", self.so_sanh, LUC,
                  280).pack(side="left", padx=10)

        ctk.CTkLabel(cha, text="Chọn hóa đơn A ở bảng trên, hóa đơn B ở bảng "
                               "dưới (bấm một dòng ở mỗi bảng), rồi bấm So sánh.",
                     text_color="#667",
                     font=ctk.CTkFont("Segoe UI", 11)).pack(anchor="w", padx=12,
                                                            pady=(6, 0))
        cot = [("ngay", "Ngày lập"), ("kh", "Ký hiệu"), ("so", "Số HĐ"),
               ("ten", "Tên người bán"), ("tien", "Tổng tiền"),
               ("tt", "Trạng thái")]
        rong = [95, 95, 95, 420, 120, 110]
        ctk.CTkLabel(cha, text="HÓA ĐƠN A",
                     font=ctk.CTkFont("Segoe UI", 12, "bold"),
                     text_color=XANH_DAM).pack(anchor="w", padx=12, pady=(8, 0))
        self.tv_a = self._bang(cha, cot, rong, 8)
        ctk.CTkLabel(cha, text="HÓA ĐƠN B",
                     font=ctk.CTkFont("Segoe UI", 12, "bold"),
                     text_color=XANH_DAM).pack(anchor="w", padx=12, pady=(8, 0))
        self.tv_b = self._bang(cha, cot, rong, 8)
        self._ban_do_ss = {}

    def loc_so_sanh(self):
        tu_khoa = dc.khong_dau(self.o_loc.get())
        self.tv_a.delete(*self.tv_a.get_children())
        self.tv_b.delete(*self.tv_b.get_children())
        self._ban_do_ss = {}
        dem = 0
        for h in self.ds:
            chuoi = dc.khong_dau("%s %s %s %s" % (h["so_hoa_don"],
                                                  h["ky_hieu_hoa_don"],
                                                  h["ten_nguoi_ban"],
                                                  h["mst_nguoi_ban"]))
            if tu_khoa and tu_khoa not in chuoi:
                continue
            gt = (dc.ngay_vn(h["ngay_lap"]), h["ky_hieu_mau_so"] +
                  h["ky_hieu_hoa_don"], h["so_hoa_don"], h["ten_nguoi_ban"],
                  dc.fmt_tien(h["tong_tien"]), h["trang_thai"])
            ia = self.tv_a.insert("", "end", values=gt)
            ib = self.tv_b.insert("", "end", values=gt)
            self._ban_do_ss[("a", ia)] = h
            self._ban_do_ss[("b", ib)] = h
            dem += 1
            if dem >= 400:
                break
        self.bao("Hiển thị %d hóa đơn để so sánh." % dem)

    def so_sanh(self):
        sa, sb = self.tv_a.selection(), self.tv_b.selection()
        if not sa or not sb:
            messagebox.showinfo("Chưa chọn đủ",
                                "Hãy bấm chọn một dòng ở bảng HÓA ĐƠN A và "
                                "một dòng ở bảng HÓA ĐƠN B.")
            return
        a = self._ban_do_ss.get(("a", sa[0]))
        b = self._ban_do_ss.get(("b", sb[0]))
        if not a or not b:
            messagebox.showinfo("Lỗi", "Không lấy được dữ liệu dòng đã chọn.")
            return
        thu_muc = os.path.join(THU_MUC, "ban_the_hien")
        os.makedirs(thu_muc, exist_ok=True)
        d = os.path.join(thu_muc, "SoSanh_%s_%s_%s.html"
                         % (a["so_hoa_don"], b["so_hoa_don"],
                            datetime.datetime.now().strftime("%H%M%S")))
        try:
            bc.so_sanh_hai_hoa_don(a, b, d)
            webbrowser.open("file:///" + d.replace("\\", "/"))
            self.bao("Đã mở bản so sánh: %s" % d)
        except Exception as e:
            messagebox.showerror("Lỗi", str(e))

    # ------------------------------------------------------- chung
    def nap_kho(self):
        try:
            self.ds = dc.doc_hoa_don()
        except Exception as e:
            messagebox.showerror("Không mở được kho dữ liệu",
                                 "%s\n\nHãy chạy ChayApp.bat ít nhất một lần "
                                 "để tạo cơ sở dữ liệu." % e)
            self.ds = []
        self.lb_kho.configure(text="Kho: %d hóa đơn" % len(self.ds))
        self.bao("Đã nạp %d hóa đơn từ %s" % (len(self.ds), dc.DUONG_DAN_DB))
        try:
            self.loc_so_sanh()
        except Exception:
            pass


def main():
    try:
        CongCu().mainloop()
    except Exception:
        thu_muc = os.path.join(THU_MUC, "logs")
        os.makedirs(thu_muc, exist_ok=True)
        with open(os.path.join(thu_muc, "app.log"), "a", encoding="utf-8") as f:
            f.write("\n%s [ERROR] cong_cu.py\n%s\n"
                    % (datetime.datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
                       traceback.format_exc()))
        raise


if __name__ == "__main__":
    main()
