# -*- coding: utf-8 -*-
"""
va_excel.py - TAX INVOICE PRO v1.7.1
Sua cong thuc cong nhom trong tep Excel xuat ra (loi SUM lech dong),
va tu gan vao cac ham xuat bao cao de tep moi luon dung ngay tu dau.
Tac gia: TDA - 0986509106
"""

import os
import sys
import glob
import functools
import unicodedata
import traceback
import datetime as dt

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

PHIEN_BAN = "v1.7.1"
THU_MUC = os.path.dirname(os.path.abspath(__file__))


# ------------------------------------------------------------------ tien ich
def _kd(s):
    """Bo dau tieng Viet, ve chu thuong de so khop nhan dong."""
    s = unicodedata.normalize("NFD", str(s))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return s.replace("\u0111", "d").replace("\u0110", "D").strip().lower()


def _la_so(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _la_ct(v):
    return isinstance(v, str) and v.startswith("=")


def _nhan_dong(nhan):
    """Xac dinh dong tieu de nhom / dong cong nhom / dong tong cong."""
    if nhan.startswith("nhom"):
        return "nhom"
    if nhan.startswith("tong cong") or nhan.startswith("tong so") or nhan == "tong":
        return "tong"
    if nhan.startswith("cong nhom") or nhan.startswith("cong theo") or nhan == "cong":
        return "cong"
    return ""


def phan_loai(ws):
    """Tra ve {so_dong: loai} voi loai thuoc trong/chu/nhom/dl/cong/tong."""
    loai = {}
    for r in range(1, (ws.max_row or 0) + 1):
        chuoi, so, ct = [], 0, 0
        for c in range(1, (ws.max_column or 0) + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            if isinstance(v, str):
                s = v.strip()
                if not s:
                    continue
                if s.startswith("="):
                    ct += 1
                else:
                    chuoi.append(_kd(s))
            elif _la_so(v):
                so += 1
        if not chuoi and not so and not ct:
            loai[r] = "trong"
            continue
        kieu = ""
        if len(chuoi) <= 3:            # dong nhan thuong rat it o chu
            for n in chuoi:
                kieu = _nhan_dong(n)
                if kieu:
                    break
        if kieu:
            loai[r] = kieu
        elif so or ct:
            loai[r] = "dl"
        else:
            loai[r] = "chu"            # tieu de bang, ghi chu
    return loai


def _co_du_lieu(ws, cot, hang):
    for h in hang:
        v = ws.cell(row=h, column=cot).value
        if _la_so(v) or _la_ct(v):
            return True
    return False


# ------------------------------------------------------------------ sua sheet
def sua_sheet(ws):
    """Tinh lai pham vi SUM cho moi dong cong nhom / tong cong. Tra ve ds sua."""
    loai = phan_loai(ws)
    cac_dong_cong = [r for r in sorted(loai) if loai[r] == "cong"]
    sua = []

    for r in sorted(loai):
        kieu = loai[r]
        if kieu not in ("cong", "tong"):
            continue

        if kieu == "cong":
            i = r - 1
            while i >= 1 and loai.get(i) == "dl":
                i -= 1
            dau, cuoi = i + 1, r - 1
            if dau > cuoi:
                continue
            hang = list(range(dau, cuoi + 1))
            lien_tuc = True
        else:
            truoc = [x for x in cac_dong_cong if x < r]
            if truoc:
                hang = truoc            # tong cong = cong cac dong nhom
            else:
                i = r - 1
                while i >= 1 and loai.get(i) in ("dl", "nhom", "chu"):
                    i -= 1
                hang = [x for x in range(i + 1, r) if loai.get(x) == "dl"]
            if not hang:
                continue
            lien_tuc = (len(hang) == hang[-1] - hang[0] + 1)

        for c in range(1, (ws.max_column or 0) + 1):
            o = ws.cell(row=r, column=c)
            hien = o.value
            if not (_la_ct(hien) or _la_so(hien)):
                continue
            if _la_ct(hien) and "SUM" not in hien.upper():
                continue
            if not _co_du_lieu(ws, c, hang):
                continue
            L = get_column_letter(c)
            if lien_tuc:
                moi = "=SUM(%s%d:%s%d)" % (L, hang[0], L, hang[-1])
            else:
                moi = "=SUM(%s)" % ",".join("%s%d" % (L, h) for h in hang)
            if str(hien) != moi:
                sua.append((ws.title, "%s%d" % (L, r), str(hien), moi))
                o.value = moi
    return sua


def sua_tep(duong_dan, in_ra=False):
    """Mo tep xlsx, sua moi cong thuc cong nhom bi lech, luu lai."""
    if not os.path.exists(duong_dan):
        return 0, []
    if not duong_dan.lower().endswith((".xlsx", ".xlsm")):
        return 0, []
    try:
        wb = load_workbook(duong_dan)   # giu nguyen cong thuc
    except Exception:
        if in_ra:
            print("[loi] khong mo duoc %s" % duong_dan)
            print(traceback.format_exc())
        return 0, []
    tat_ca = []
    for ws in wb.worksheets:
        try:
            tat_ca += sua_sheet(ws)
        except Exception:
            if in_ra:
                print("[loi] sheet %s: %s" % (ws.title, traceback.format_exc()))
    if tat_ca:
        try:
            wb.save(duong_dan)
        except Exception:
            if in_ra:
                print("[loi] khong luu duoc (tep dang mo trong Excel?): %s" % duong_dan)
            try:
                wb.close()
            except Exception:
                pass
            return 0, tat_ca
    try:
        wb.close()
    except Exception:
        pass
    if in_ra:
        print("%s -> sua %d cong thuc" % (os.path.basename(duong_dan), len(tat_ca)))
        for sheet, o, cu, moi in tat_ca[:60]:
            print("   [%s] %s : %s  ->  %s" % (sheet, o, cu, moi))
        if len(tat_ca) > 60:
            print("   ... va %d cong thuc khac" % (len(tat_ca) - 60))
    return len(tat_ca), tat_ca


def sua_thu_muc(thu_muc, de_quy=True, in_ra=True):
    mau = "**/*.xlsx" if de_quy else "*.xlsx"
    tong, tep = 0, 0
    for dd in glob.glob(os.path.join(thu_muc, mau), recursive=de_quy):
        if os.path.basename(dd).startswith("~$"):
            continue
        n, _ = sua_tep(dd, in_ra=in_ra)
        tong += n
        tep += 1
    if in_ra:
        print("Da kiem tra %d tep, sua tong cong %d cong thuc." % (tep, tong))
    return tong


# ------------------------------------------------------------------ tu dong va
def _tim_xlsx(gt):
    ds = []
    if isinstance(gt, str):
        ds.append(gt)
    elif isinstance(gt, dict):
        for v in gt.values():
            ds += _tim_xlsx(v)
    elif isinstance(gt, (list, tuple, set)):
        for v in gt:
            ds += _tim_xlsx(v)
    return [d for d in ds
            if isinstance(d, str) and d.lower().endswith((".xlsx", ".xlsm")) and os.path.exists(d)]


def _boc(ham):
    @functools.wraps(ham)
    def chay(*a, **kw):
        kq = ham(*a, **kw)
        try:
            dd = _tim_xlsx(kq) or _tim_xlsx(a) or _tim_xlsx(kw)
            for d in dict.fromkeys(dd):
                sua_tep(d)
        except Exception:
            pass
        return kq
    chay._tda_boc = True
    return chay


def va_module(ten):
    """Boc moi ham xuat_* cua module de tu sua cong thuc sau khi xuat."""
    try:
        m = __import__(ten)
    except Exception:
        return 0
    if getattr(m, "_tda_va_excel", False):
        return 0
    n = 0
    for ten_ham in dir(m):
        if not ten_ham.startswith("xuat"):
            continue
        ham = getattr(m, ten_ham, None)
        if not callable(ham) or isinstance(ham, type):
            continue
        if getattr(ham, "_tda_boc", False):
            continue
        try:
            setattr(m, ten_ham, _boc(ham))
            n += 1
        except Exception:
            pass
    m._tda_va_excel = True
    return n


def va_tat_ca():
    tong = 0
    for ten in ("bao_cao_chi_tiet", "bao_cao", "exporter"):
        tong += va_module(ten)
    return tong


# ------------------------------------------------------------------ CLI
def main():
    print("=== VA CONG THUC EXCEL - TAX INVOICE PRO %s ===" % PHIEN_BAN)
    print("Thoi diem: %s" % dt.datetime.now().strftime("%d/%m/%Y %H:%M:%S"))
    dich = sys.argv[1] if len(sys.argv) > 1 else THU_MUC
    if os.path.isdir(dich):
        print("Quet thu muc: %s" % dich)
        sua_thu_muc(dich, de_quy=True, in_ra=True)
    else:
        n, _ = sua_tep(dich, in_ra=True)
        print("Xong, sua %d cong thuc." % n)
    print("Mo lai tep bang Excel de kiem tra dong 'Cong nhom'.")


if __name__ == "__main__":
    main()


TEP_DAY_DU = True  # dau hieu tep da duoc dan day du - dung xoa dong nay
