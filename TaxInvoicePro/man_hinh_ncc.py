# -*- coding: utf-8 -*-
"""
man_hinh_ncc.py - TAX INVOICE PRO v1.6.1
Man hinh "Nha cung cap & Rui ro" (3 tab, moi tab boc rieng chong loi).
Tac gia: TDA - 0986509106
Luu y: moi tra cuu truc tuyen deu MAC DINH TAT, nguoi dung tu bat trong tab Cau hinh.
"""

import os
import re
import csv
import json
import time
import sqlite3
import threading
import traceback
import webbrowser
import datetime as dt
import urllib.request

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

try:
    import va_ttk  # va loi command=None cua ttk tren Python 3.12+
except Exception:
    pass

import customtkinter as ctk

# ----------------------------------------------------------------------------
# Hang so
# ----------------------------------------------------------------------------
PHIEN_BAN = "v1.6.1"
TAC_GIA = "TDA - 0986509106"

THU_MUC = os.path.dirname(os.path.abspath(__file__))
THU_MUC_DL = os.path.join(THU_MUC, "data")
os.makedirs(THU_MUC_DL, exist_ok=True)

TEP_CACHE = os.path.join(THU_MUC_DL, "mst_cache.json")
TEP_RUI_RO = os.path.join(THU_MUC_DL, "ds_rui_ro.json")
TEP_CAU_HINH = os.path.join(THU_MUC_DL, "mst_cau_hinh.json")
TEP_DB = os.path.join(THU_MUC_DL, "taxinvoice.db")

XANH = "#1565C0"
XANH_DAM = "#0D47A1"
NEN = "#F4F7FB"
TRANG = "#FFFFFF"
VIEN = "#CBD5E1"
CAM = "#EF6C00"
DO = "#C62828"
LUC = "#2E7D32"
XAM = "#64748B"

URL_MAC_DINH = "https://api.vietqr.io/v2/business/{mst}"
URL_DS_RUI_RO = "https://tracuuhoadon.gdt.gov.vn/dnrrvp.html"
URL_TRA_TAY = "https://tracuunnt.gdt.gov.vn/tcnnt/mstdn.jsp"

# ----------------------------------------------------------------------------
# Nap module phu (khong bat buoc)
# ----------------------------------------------------------------------------
LOI_MST = ""
try:
    import mst_tra_cuu as mst
except Exception:
    mst = None
    LOI_MST = traceback.format_exc()

try:
    import kho
except Exception:
    kho = None


def _hs(ten, mac_dinh):
    """Lay hang so tu mst_tra_cuu neu co."""
    if mst is not None:
        gt = getattr(mst, ten, None)
        if gt:
            return gt
    return mac_dinh


RUI_RO = _hs("RUI_RO", "RUI_RO")
THEO_DOI = _hs("THEO_DOI", "THEO_DOI")
AN_TOAN = _hs("AN_TOAN", "AN_TOAN")
KHONG_RO = _hs("KHONG_RO", "KHONG_RO")

NHAN_MUC = _hs("NHAN_MUC", None) or {}
NHAN_MUC.setdefault(RUI_RO, "Rủi ro")
NHAN_MUC.setdefault(THEO_DOI, "Cần theo dõi")
NHAN_MUC.setdefault(AN_TOAN, "Bình thường")
NHAN_MUC.setdefault(KHONG_RO, "Chưa tra cứu")

MAU_MUC = {
    RUI_RO: ("#FFEBEE", DO),
    THEO_DOI: ("#FFF8E1", CAM),
    AN_TOAN: ("#E8F5E9", LUC),
    KHONG_RO: (TRANG, "#1F2937"),
}

CAU_HINH_MAC_DINH = {
    "bat": False,
    "url_mau": URL_MAC_DINH,
    "so_ngay_cache": 7,
    "do_tre": 0.6,
    "thoi_gian_cho": 10,
}

# ----------------------------------------------------------------------------
# Tien ich chung
# ----------------------------------------------------------------------------


def _fmt_tien(v):
    try:
        return "{:,.0f}".format(float(v or 0)).replace(",", ".")
    except Exception:
        return "0"


def _bay_gio():
    return dt.datetime.now().strftime("%d/%m/%Y %H:%M")


def _doc_json(duong_dan, mac_dinh):
    try:
        if os.path.exists(duong_dan):
            with open(duong_dan, "r", encoding="utf-8") as f:
                return json.load(f)
    except Exception:
        pass
    return mac_dinh


def _ghi_json(duong_dan, du_lieu):
    try:
        tam = duong_dan + ".tmp"
        with open(tam, "w", encoding="utf-8") as f:
            json.dump(du_lieu, f, ensure_ascii=False, indent=2)
        if os.path.exists(duong_dan):
            os.remove(duong_dan)
        os.rename(tam, duong_dan)
        return True
    except Exception:
        return False


def doc_cau_hinh():
    cfg = dict(CAU_HINH_MAC_DINH)
    if mst is not None and hasattr(mst, "doc_cau_hinh"):
        try:
            tu_module = mst.doc_cau_hinh() or {}
            if isinstance(tu_module, dict):
                cfg.update(tu_module)
                return cfg
        except Exception:
            pass
    cfg.update(_doc_json(TEP_CAU_HINH, {}) or {})
    return cfg


def ghi_cau_hinh(cfg):
    if mst is not None and hasattr(mst, "ghi_cau_hinh"):
        try:
            mst.ghi_cau_hinh(cfg)
        except Exception:
            pass
    return _ghi_json(TEP_CAU_HINH, cfg)


def doc_cache():
    """Tra ve dict {mst: ket_qua}."""
    kq = {}
    if mst is not None and hasattr(mst, "doc_cache_tat_ca"):
        try:
            dl = mst.doc_cache_tat_ca()
            if isinstance(dl, dict):
                kq.update(dl)
            elif isinstance(dl, (list, tuple)):
                for d in dl:
                    if isinstance(d, dict) and d.get("mst"):
                        kq[str(d["mst"])] = d
        except Exception:
            pass
    for k, v in (_doc_json(TEP_CACHE, {}) or {}).items():
        kq.setdefault(str(k), v)
    return kq


def ghi_cache(cache):
    return _ghi_json(TEP_CACHE, cache)


def doc_ds_rui_ro():
    """Tra ve list dict {mst, ten, ly_do, nguon, ngay}."""
    ds = []
    if mst is not None and hasattr(mst, "doc_ds_rui_ro"):
        try:
            dl = mst.doc_ds_rui_ro() or []
            for d in dl:
                if isinstance(d, dict):
                    ds.append(d)
                elif d:
                    ds.append({"mst": str(d)})
        except Exception:
            pass
    if not ds:
        ds = _doc_json(TEP_RUI_RO, []) or []
    chuan = []
    for d in ds:
        if not isinstance(d, dict):
            d = {"mst": str(d)}
        m = re.sub(r"\D", "", str(d.get("mst", "")))
        if not m:
            continue
        chuan.append({
            "mst": m,
            "ten": d.get("ten", "") or d.get("ten_ncc", ""),
            "ly_do": d.get("ly_do", "") or d.get("ghi_chu", ""),
            "nguon": d.get("nguon", "Tệp nạp tay"),
            "ngay": d.get("ngay", ""),
        })
    return chuan


def ghi_ds_rui_ro(ds):
    return _ghi_json(TEP_RUI_RO, ds)


# ----------------------------------------------------------------------------
# Lay danh sach nha cung cap tu kho du lieu
# ----------------------------------------------------------------------------


def _gop_ncc(dong_hd):
    gom = {}
    for hd in dong_hd:
        m = re.sub(r"\D", "", str(hd.get("mst_nguoi_ban", "") or ""))
        if not m:
            continue
        g = gom.setdefault(m, {"mst": m, "ten": "", "so_hd": 0, "tong_tien": 0})
        ten = (hd.get("ten_nguoi_ban") or "").strip()
        if ten and len(ten) > len(g["ten"]):
            g["ten"] = ten
        g["so_hd"] += 1
        try:
            g["tong_tien"] += float(hd.get("tong_tien") or 0)
        except Exception:
            pass
    return sorted(gom.values(), key=lambda x: -x["tong_tien"])


def _ncc_tu_sqlite():
    if not os.path.exists(TEP_DB):
        return []
    cn = sqlite3.connect(TEP_DB)
    try:
        cn.row_factory = sqlite3.Row
        c = cn.cursor()
        c.execute("""SELECT mst_nguoi_ban AS mst,
                            MAX(ten_nguoi_ban) AS ten,
                            COUNT(*) AS so_hd,
                            SUM(tong_tien) AS tong_tien
                       FROM hoa_don
                      WHERE mst_nguoi_ban IS NOT NULL AND mst_nguoi_ban <> ''
                      GROUP BY mst_nguoi_ban""")
        ds = []
        for r in c.fetchall():
            ds.append({
                "mst": re.sub(r"\D", "", str(r["mst"] or "")),
                "ten": r["ten"] or "",
                "so_hd": int(r["so_hd"] or 0),
                "tong_tien": float(r["tong_tien"] or 0),
            })
        return [d for d in ds if d["mst"]]
    finally:
        cn.close()


def lay_ncc_tu_kho():
    """Thu lan luot: mst_tra_cuu -> kho -> SQLite truc tiep."""
    if mst is not None and hasattr(mst, "tong_hop_ncc"):
        try:
            dl = mst.tong_hop_ncc() or []
            ds = []
            for d in dl:
                if isinstance(d, dict) and d.get("mst"):
                    ds.append({
                        "mst": re.sub(r"\D", "", str(d.get("mst"))),
                        "ten": d.get("ten") or d.get("ten_ncc") or d.get("ten_nguoi_ban") or "",
                        "so_hd": int(d.get("so_hd") or d.get("so_hoa_don") or 0),
                        "tong_tien": float(d.get("tong_tien") or 0),
                    })
            if ds:
                return sorted(ds, key=lambda x: -x["tong_tien"])
        except Exception:
            pass
    if kho is not None and hasattr(kho, "lay_danh_sach"):
        try:
            dl = kho.lay_danh_sach()
            if dl:
                return _gop_ncc([dict(x) for x in dl])
        except Exception:
            pass
    return _ncc_tu_sqlite()


# ----------------------------------------------------------------------------
# Tra cuu mot ma so thue
# ----------------------------------------------------------------------------

_TU_XAU = ("ngừng", "ngung", "tạm nghỉ", "tam nghi", "không hoạt động",
           "khong hoat dong", "bỏ địa chỉ", "bo dia chi", "chờ làm thủ tục",
           "cho lam thu tuc", "đóng mã", "dong ma", "giải thể", "giai the",
           "phá sản", "pha san", "rủi ro", "rui ro")
_TU_TOT = ("đang hoạt động", "dang hoat dong", "nnt đang hoạt động")


def _suy_muc(tinh_trang):
    t = (tinh_trang or "").lower()
    if not t:
        return KHONG_RO
    for k in _TU_XAU:
        if k in t:
            return RUI_RO
    for k in _TU_TOT:
        if k in t:
            return AN_TOAN
    return THEO_DOI


def _chuan_ket_qua(ma_so, dl):
    """Chuan hoa ket qua tra cuu ve mot dict thong nhat."""
    kq = {"mst": ma_so, "ten": "", "dia_chi": "", "tinh_trang": "",
          "nguon": "", "cap_nhat": _bay_gio(), "muc": KHONG_RO, "tho": None}
    if dl is None:
        kq["tinh_trang"] = "Không lấy được dữ liệu"
        return kq
    if not isinstance(dl, dict):
        kq["tinh_trang"] = str(dl)
        kq["muc"] = _suy_muc(kq["tinh_trang"])
        return kq

    goc = dl.get("data") if isinstance(dl.get("data"), dict) else dl
    kq["tho"] = dl

    for k in ("ten", "ten_ncc", "name", "tenNNT", "ten_nnt", "shortName", "tenChinhThuc"):
        if goc.get(k):
            kq["ten"] = str(goc[k])
            break
    for k in ("dia_chi", "address", "diaChi", "dchiNNT"):
        if goc.get(k):
            kq["dia_chi"] = str(goc[k])
            break
    for k in ("tinh_trang", "status", "tthai", "trang_thai", "tinhTrang", "ttnnt"):
        if goc.get(k):
            kq["tinh_trang"] = str(goc[k])
            break
    for k in ("nguon", "source"):
        if goc.get(k):
            kq["nguon"] = str(goc[k])
            break
    if not kq["nguon"]:
        kq["nguon"] = "API tra cứu MST"

    muc = goc.get("muc") or dl.get("muc")
    kq["muc"] = muc if muc in MAU_MUC else _suy_muc(kq["tinh_trang"])
    if not kq["tinh_trang"] and kq["ten"]:
        kq["tinh_trang"] = "Có dữ liệu, không ghi tình trạng"
        kq["muc"] = THEO_DOI
    return kq


def _goi_api(ma_so, cfg):
    url = (cfg.get("url_mau") or URL_MAC_DINH)
    if "{mst}" not in url:
        raise ValueError("Mẫu địa chỉ phải chứa {mst}")
    url = url.replace("{mst}", ma_so)
    req = urllib.request.Request(url, headers={
        "User-Agent": "TaxInvoicePro/1.6.1 (desktop)",
        "Accept": "application/json",
    })
    cho = float(cfg.get("thoi_gian_cho") or 10)
    with urllib.request.urlopen(req, timeout=cho) as r:
        raw = r.read().decode("utf-8", "ignore")
    try:
        return json.loads(raw)
    except Exception:
        return {"tinh_trang": raw[:200]}


def tra_cuu_mot(ma_so, cfg):
    """Uu tien ham cua mst_tra_cuu, khong co thi goi API truc tiep."""
    if mst is not None:
        for ten in ("tra_cuu_mot", "tra_cuu_mst", "tra_cuu", "tra_cuu_mot_mst"):
            ham = getattr(mst, ten, None)
            if not callable(ham):
                continue
            for doi_so in ((ma_so, cfg), (ma_so,)):
                try:
                    return _chuan_ket_qua(ma_so, ham(*doi_so))
                except TypeError:
                    continue
                except Exception:
                    break
    try:
        return _chuan_ket_qua(ma_so, _goi_api(ma_so, cfg))
    except Exception as e:
        kq = _chuan_ket_qua(ma_so, None)
        kq["tinh_trang"] = "Lỗi: %s" % e
        return kq


# ----------------------------------------------------------------------------
# Man hinh chinh
# ----------------------------------------------------------------------------


class ManHinhNCC(ctk.CTkFrame):
    """Man hinh menu trai: Nha cung cap & Rui ro."""

    def __init__(self, master, app=None, **kw):
        super().__init__(master, fg_color=NEN, **kw)
        self.app = app
        self.cau_hinh = doc_cau_hinh()
        self.cache = doc_cache()
        self.ds_rui_ro = doc_ds_rui_ro()
        self.map_rui_ro = {d["mst"]: d for d in self.ds_rui_ro}
        self.ds_ncc = []
        self.dang_chay = False
        self.yeu_cau_dung = False
        self.tv = None
        self.tv_rr = None

        self._kieu_bang()
        self._dung_tieu_de()

        self.tabs = ctk.CTkTabview(self, fg_color=TRANG,
                                   segmented_button_selected_color=XANH,
                                   segmented_button_selected_hover_color=XANH_DAM)
        self.tabs.pack(fill="both", expand=True, padx=14, pady=(0, 6))
        self._boc("Tra cứu tình trạng NCC", self._tab_tra_cuu)
        self._boc("Danh sách rủi ro", self._tab_rui_ro)
        self._boc("Hướng dẫn & Cấu hình", self._tab_cau_hinh)

        ctk.CTkLabel(self, text="TAX INVOICE PRO %s  •  Tác giả: %s  •  Dữ liệu tra cứu chỉ mang tính tham khảo"
                     % (PHIEN_BAN, TAC_GIA),
                     text_color=XAM, font=("Segoe UI", 11)).pack(pady=(0, 8))

        self.after(250, self.nap_tu_kho)

    # -- khung chung ---------------------------------------------------------

    def _kieu_bang(self):
        st = ttk.Style()
        try:
            st.theme_use("clam")
        except Exception:
            pass
        st.configure("NCC.Treeview", background=TRANG, fieldbackground=TRANG,
                     foreground="#1F2937", rowheight=27, font=("Segoe UI", 10),
                     bordercolor=VIEN, borderwidth=0)
        st.configure("NCC.Treeview.Heading", background="#E3F0FD",
                     foreground=XANH_DAM, font=("Segoe UI", 10, "bold"))
        st.map("NCC.Treeview", background=[("selected", "#BBDEFB")],
               foreground=[("selected", "#0B2545")])

    def _dung_tieu_de(self):
        d = ctk.CTkFrame(self, fg_color="transparent")
        d.pack(fill="x", padx=14, pady=(12, 8))
        ctk.CTkLabel(d, text="NHÀ CUNG CẤP & CẢNH BÁO RỦI RO",
                     font=("Segoe UI", 18, "bold"), text_color=XANH_DAM).pack(anchor="w")
        ctk.CTkLabel(d, text="Tổng hợp người bán từ kho hóa đơn, đối chiếu tình trạng mã số thuế "
                             "và danh sách doanh nghiệp rủi ro. Mặc định không kết nối ra ngoài.",
                     font=("Segoe UI", 12), text_color=XAM).pack(anchor="w", pady=(2, 0))

    def _boc(self, ten, ham):
        tab = self.tabs.add(ten)
        try:
            ham(tab)
        except Exception:
            self._khung_loi(tab, traceback.format_exc())

    def _khung_loi(self, cha, noi_dung):
        ctk.CTkLabel(cha, text="Tab này gặp lỗi, các tab khác vẫn dùng được:",
                     text_color=DO, font=("Segoe UI", 12, "bold")).pack(anchor="w", padx=10, pady=(10, 4))
        tb = ctk.CTkTextbox(cha, height=260, font=("Consolas", 10))
        tb.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        tb.insert("1.0", noi_dung)
        tb.configure(state="disabled")

    def _bang(self, cha, cot, ten_tv):
        khung = tk.Frame(cha, bg=TRANG, highlightthickness=1, highlightbackground=VIEN)
        khung.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        tv = ttk.Treeview(khung, columns=[c[0] for c in cot], show="headings",
                          style="NCC.Treeview", selectmode="extended")
        for ma, nhan, rong, neo in cot:
            tv.heading(ma, text=nhan, command=lambda m=ma, t=ten_tv: self._sap_xep(t, m))
            tv.column(ma, width=rong, anchor=neo, stretch=(neo == "w" and rong > 180))
        vs = ttk.Scrollbar(khung, orient="vertical", command=tv.yview)
        hs = ttk.Scrollbar(khung, orient="horizontal", command=tv.xview)
        tv.configure(yscrollcommand=vs.set, xscrollcommand=hs.set)
        tv.grid(row=0, column=0, sticky="nsew")
        vs.grid(row=0, column=1, sticky="ns")
        hs.grid(row=1, column=0, sticky="ew")
        khung.rowconfigure(0, weight=1)
        khung.columnconfigure(0, weight=1)
        for ma_muc, (nen, chu) in MAU_MUC.items():
            tv.tag_configure(ma_muc, background=nen, foreground=chu)
        return tv

    def _sap_xep(self, ten_tv, cot):
        tv = getattr(self, ten_tv, None)
        if tv is None:
            return
        dong = [(tv.set(k, cot), k) for k in tv.get_children("")]

        def khoa(x):
            v = (x[0] or "").replace(".", "").replace(" ", "")
            try:
                return (0, float(v))
            except Exception:
                return (1, (x[0] or "").lower())

        nguoc = getattr(self, "_nguoc_" + cot, False)
        dong.sort(key=khoa, reverse=nguoc)
        setattr(self, "_nguoc_" + cot, not nguoc)
        for i, (_, k) in enumerate(dong):
            tv.move(k, "", i)
        self._danh_so()

    def _danh_so(self):
        if self.tv is None:
            return
        for i, k in enumerate(self.tv.get_children(""), 1):
            self.tv.set(k, "stt", str(i))

    # ========================= TAB 1: TRA CUU ===============================

    def _tab_tra_cuu(self, cha):
        thanh = ctk.CTkFrame(cha, fg_color="transparent")
        thanh.pack(fill="x", padx=10, pady=10)

        ctk.CTkButton(thanh, text="Nạp danh sách NCC từ kho", width=190,
                      fg_color=XANH, hover_color=XANH_DAM,
                      command=self.nap_tu_kho).pack(side="left")
        ctk.CTkButton(thanh, text="TRA CỨU DÒNG ĐANG CHỌN", width=200,
                      fg_color=XANH_DAM, hover_color=XANH,
                      command=lambda: self.tra_cuu(True)).pack(side="left", padx=6)
        ctk.CTkButton(thanh, text="Tra cứu tất cả chưa có kết quả", width=215,
                      fg_color="#37474F", hover_color="#263238",
                      command=lambda: self.tra_cuu(False)).pack(side="left")
        self.nut_dung = ctk.CTkButton(thanh, text="Dừng", width=70, state="disabled",
                                      fg_color=DO, hover_color="#8E0000",
                                      command=self.dung_tra_cuu)
        self.nut_dung.pack(side="left", padx=6)
        ctk.CTkButton(thanh, text="Xuất CSV", width=90, fg_color="#455A64",
                      hover_color="#263238", command=self.xuat_csv).pack(side="right")

        loc = ctk.CTkFrame(cha, fg_color="transparent")
        loc.pack(fill="x", padx=10, pady=(0, 8))
        ctk.CTkLabel(loc, text="Tìm theo MST / tên:", font=("Segoe UI", 12)).pack(side="left")
        self.o_tim = ctk.CTkEntry(loc, width=260, placeholder_text="Gõ rồi nhấn Enter")
        self.o_tim.pack(side="left", padx=8)
        self.o_tim.bind("<Return>", lambda e: self.ve_bang())
        ctk.CTkLabel(loc, text="Mức:", font=("Segoe UI", 12)).pack(side="left", padx=(10, 4))
        self.chon_muc = ctk.CTkOptionMenu(
            loc, width=170, values=["Tất cả", NHAN_MUC[RUI_RO], NHAN_MUC[THEO_DOI],
                                    NHAN_MUC[AN_TOAN], NHAN_MUC[KHONG_RO]],
            fg_color=XANH, button_color=XANH_DAM, command=lambda v: self.ve_bang())
        self.chon_muc.pack(side="left")
        ctk.CTkButton(loc, text="Xóa bộ lọc", width=90, fg_color="#78909C",
                      hover_color="#546E7A", command=self.xoa_loc).pack(side="left", padx=8)

        cot = [("stt", "STT", 50, "center"),
               ("mst", "Mã số thuế", 130, "center"),
               ("ten", "Tên nhà cung cấp", 330, "w"),
               ("so_hd", "Số HĐ", 70, "center"),
               ("tong_tien", "Tổng tiền (đ)", 130, "e"),
               ("tinh_trang", "Tình trạng", 230, "w"),
               ("muc", "Đánh giá", 120, "center"),
               ("nguon", "Nguồn", 150, "w"),
               ("cap_nhat", "Cập nhật", 120, "center")]
        self.tv = self._bang(cha, cot, "tv")
        self.tv.bind("<Double-1>", self.xem_chi_tiet)

        duoi = ctk.CTkFrame(cha, fg_color="transparent")
        duoi.pack(fill="x", padx=10, pady=(0, 10))
        self.thanh_tt = ctk.CTkProgressBar(duoi, height=8, progress_color=XANH)
        self.thanh_tt.pack(fill="x", pady=(0, 6))
        self.thanh_tt.set(0)
        self.nhan_tt = ctk.CTkLabel(duoi, text="Sẵn sàng.", font=("Segoe UI", 11), text_color=XAM)
        self.nhan_tt.pack(anchor="w")

    def xoa_loc(self):
        self.o_tim.delete(0, "end")
        self.chon_muc.set("Tất cả")
        self.ve_bang()

    def nap_tu_kho(self):
        try:
            self.ds_ncc = lay_ncc_tu_kho()
        except Exception:
            self.ds_ncc = []
            messagebox.showerror("Lỗi", "Không đọc được kho hóa đơn:\n\n" + traceback.format_exc())
        self.ve_bang()

    def _muc_cua(self, ma_so):
        if ma_so in self.map_rui_ro:
            return RUI_RO
        kq = self.cache.get(ma_so)
        if not kq:
            return KHONG_RO
        return kq.get("muc", KHONG_RO)

    def ve_bang(self):
        if self.tv is None:
            return
        tim = (self.o_tim.get() or "").strip().lower()
        chon = self.chon_muc.get()
        for k in self.tv.get_children(""):
            self.tv.delete(k)
        dem = {RUI_RO: 0, THEO_DOI: 0, AN_TOAN: 0, KHONG_RO: 0}
        stt = 0
        for n in self.ds_ncc:
            ma_so = n["mst"]
            muc = self._muc_cua(ma_so)
            dem[muc] = dem.get(muc, 0) + 1
            if chon != "Tất cả" and NHAN_MUC.get(muc, "") != chon:
                continue
            if tim and tim not in ma_so.lower() and tim not in (n.get("ten", "") or "").lower():
                continue
            kq = self.cache.get(ma_so, {})
            tt = kq.get("tinh_trang", "")
            if ma_so in self.map_rui_ro:
                ly_do = self.map_rui_ro[ma_so].get("ly_do") or "Có trong danh sách rủi ro"
                tt = (tt + " | " if tt else "") + ly_do
            stt += 1
            self.tv.insert("", "end", iid=ma_so, tags=(muc,), values=(
                stt, ma_so, n.get("ten", "") or kq.get("ten", ""),
                n.get("so_hd", 0), _fmt_tien(n.get("tong_tien", 0)),
                tt, NHAN_MUC.get(muc, muc), kq.get("nguon", ""), kq.get("cap_nhat", "")))
        self.nhan_tt.configure(
            text="Hiển thị %d / %d nhà cung cấp  •  Rủi ro: %d  •  Theo dõi: %d  •  Bình thường: %d  •  Chưa tra cứu: %d"
                 % (stt, len(self.ds_ncc), dem[RUI_RO], dem[THEO_DOI], dem[AN_TOAN], dem[KHONG_RO]))

    def tra_cuu(self, chi_chon):
        if self.dang_chay:
            messagebox.showinfo("Đang chạy", "Đang tra cứu, vui lòng đợi hoặc bấm Dừng.")
            return
        if not self.cau_hinh.get("bat", False):
            messagebox.showwarning(
                "Tính năng đang tắt",
                "Tra cứu trực tuyến mặc định TẮT.\n\n"
                "Mở tab 'Hướng dẫn & Cấu hình', đánh dấu 'Cho phép tra cứu trực tuyến', "
                "bấm Lưu cấu hình rồi thử lại.")
            return
        if chi_chon:
            ds = list(self.tv.selection())
            if not ds:
                messagebox.showinfo("Chưa chọn", "Hãy chọn ít nhất một dòng trong bảng.")
                return
        else:
            han = int(self.cau_hinh.get("so_ngay_cache") or 7)
            ds = [n["mst"] for n in self.ds_ncc if self._het_han(n["mst"], han)]
            if not ds:
                messagebox.showinfo("Không có gì để tra", "Mọi nhà cung cấp đều đã có kết quả còn hạn.")
                return
            if not messagebox.askyesno("Xác nhận", "Sẽ tra cứu %d mã số thuế. Tiếp tục?" % len(ds)):
                return
        self.dang_chay = True
        self.yeu_cau_dung = False
        self.nut_dung.configure(state="normal")
        self.thanh_tt.set(0)
        threading.Thread(target=self._chay, args=(ds,), daemon=True).start()

    def _het_han(self, ma_so, so_ngay):
        kq = self.cache.get(ma_so)
        if not kq:
            return True
        try:
            t = dt.datetime.strptime(kq.get("cap_nhat", ""), "%d/%m/%Y %H:%M")
            return (dt.datetime.now() - t).days >= so_ngay
        except Exception:
            return True

    def _chay(self, ds):
        tong = len(ds)
        tre = float(self.cau_hinh.get("do_tre") or 0.6)
        for i, ma_so in enumerate(ds, 1):
            if self.yeu_cau_dung:
                break
            kq = tra_cuu_mot(ma_so, self.cau_hinh)
            self.cache[ma_so] = kq
            self.after(0, self._bao_tien_do, i, tong, ma_so, kq)
            if i < tong:
                time.sleep(max(0.0, tre))
        self.after(0, self._xong)

    def _bao_tien_do(self, i, tong, ma_so, kq):
        self.thanh_tt.set(i / float(max(1, tong)))
        self.nhan_tt.configure(text="Đang tra cứu %d/%d — %s: %s"
                               % (i, tong, ma_so, kq.get("tinh_trang", "")))

    def _xong(self):
        self.dang_chay = False
        self.nut_dung.configure(state="disabled")
        ghi_cache(self.cache)
        self.thanh_tt.set(1 if not self.yeu_cau_dung else 0)
        self.ve_bang()
        messagebox.showinfo("Hoàn tất", "Đã %s tra cứu. Kết quả được lưu lại để lần sau dùng ngay."
                            % ("dừng" if self.yeu_cau_dung else "xong"))

    def dung_tra_cuu(self):
        self.yeu_cau_dung = True
        self.nhan_tt.configure(text="Đang dừng…")

    def xem_chi_tiet(self, _e=None):
        chon = self.tv.selection()
        if not chon:
            return
        ma_so = chon[0]
        kq = self.cache.get(ma_so, {})
        ncc = next((n for n in self.ds_ncc if n["mst"] == ma_so), {})
        cs = ctk.CTkToplevel(self)
        cs.title("Chi tiết nhà cung cấp %s" % ma_so)
        cs.geometry("720x520")
        cs.transient(self.winfo_toplevel())
        tb = ctk.CTkTextbox(cs, font=("Consolas", 11))
        tb.pack(fill="both", expand=True, padx=12, pady=12)
        dong = [
            "MÃ SỐ THUẾ        : %s" % ma_so,
            "TÊN (theo hóa đơn): %s" % (ncc.get("ten") or "(không có)"),
            "TÊN (theo tra cứu): %s" % (kq.get("ten") or "(chưa tra cứu)"),
            "ĐỊA CHỈ           : %s" % (kq.get("dia_chi") or "(không có)"),
            "SỐ HÓA ĐƠN        : %s" % ncc.get("so_hd", 0),
            "TỔNG TIỀN         : %s đ" % _fmt_tien(ncc.get("tong_tien", 0)),
            "TÌNH TRẠNG        : %s" % (kq.get("tinh_trang") or "(chưa tra cứu)"),
            "ĐÁNH GIÁ          : %s" % NHAN_MUC.get(self._muc_cua(ma_so), ""),
            "NGUỒN             : %s" % (kq.get("nguon") or ""),
            "CẬP NHẬT          : %s" % (kq.get("cap_nhat") or ""),
            "",
        ]
        if ma_so in self.map_rui_ro:
            d = self.map_rui_ro[ma_so]
            dong += ["CÓ TRONG DANH SÁCH RỦI RO",
                     "  Lý do : %s" % (d.get("ly_do") or "(không ghi)"),
                     "  Nguồn : %s" % (d.get("nguon") or ""),
                     "  Ngày  : %s" % (d.get("ngay") or ""), ""]
        if kq.get("tho"):
            dong += ["DỮ LIỆU GỐC TRẢ VỀ:",
                     json.dumps(kq["tho"], ensure_ascii=False, indent=2)]
        tb.insert("1.0", "\n".join(dong))
        tb.configure(state="disabled")
        hang = ctk.CTkFrame(cs, fg_color="transparent")
        hang.pack(fill="x", padx=12, pady=(0, 12))
        ctk.CTkButton(hang, text="Mở trang tra cứu tay của Cục Thuế", fg_color=XANH,
                      hover_color=XANH_DAM,
                      command=lambda: webbrowser.open(URL_TRA_TAY)).pack(side="left")
        ctk.CTkButton(hang, text="Đóng", width=90, fg_color="#607D8B",
                      hover_color="#455A64", command=cs.destroy).pack(side="right")

    def xuat_csv(self):
        if not self.ds_ncc:
            messagebox.showinfo("Chưa có dữ liệu", "Hãy nạp danh sách nhà cung cấp trước.")
            return
        dd = filedialog.asksaveasfilename(
            title="Lưu danh sách nhà cung cấp",
            defaultextension=".csv", filetypes=[("Tệp CSV", "*.csv")],
            initialfile="NhaCungCap_%s.csv" % dt.datetime.now().strftime("%Y%m%d_%H%M"))
        if not dd:
            return
        try:
            with open(dd, "w", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f, delimiter=";")
                w.writerow(["Mã số thuế", "Tên nhà cung cấp", "Số hóa đơn", "Tổng tiền",
                            "Tình trạng", "Đánh giá", "Nguồn", "Cập nhật", "Trong DS rủi ro"])
                for n in self.ds_ncc:
                    ma_so = n["mst"]
                    kq = self.cache.get(ma_so, {})
                    w.writerow([ma_so, n.get("ten", ""), n.get("so_hd", 0),
                                int(n.get("tong_tien", 0) or 0),
                                kq.get("tinh_trang", ""),
                                NHAN_MUC.get(self._muc_cua(ma_so), ""),
                                kq.get("nguon", ""), kq.get("cap_nhat", ""),
                                "Có" if ma_so in self.map_rui_ro else "Không"])
            messagebox.showinfo("Đã lưu", "Đã lưu tệp:\n%s" % dd)
        except Exception:
            messagebox.showerror("Lỗi", traceback.format_exc())

    # ========================= TAB 2: DANH SACH RUI RO ======================

    def _tab_rui_ro(self, cha):
        ctk.CTkLabel(cha,
                     text="Danh sách doanh nghiệp rủi ro do cơ quan thuế công bố. "
                          "Ứng dụng không tự tải; bạn tải tệp từ trang Cục Thuế rồi nạp vào đây.",
                     font=("Segoe UI", 12), text_color=XAM,
                     wraplength=1000, justify="left").pack(anchor="w", padx=10, pady=(10, 6))

        thanh = ctk.CTkFrame(cha, fg_color="transparent")
        thanh.pack(fill="x", padx=10, pady=(0, 8))
        ctk.CTkButton(thanh, text="Mở trang Cục Thuế để tải tệp", width=210,
                      fg_color=XANH, hover_color=XANH_DAM,
                      command=lambda: webbrowser.open(URL_DS_RUI_RO)).pack(side="left")
        ctk.CTkButton(thanh, text="Nạp tệp Excel/CSV", width=150, fg_color=XANH_DAM,
                      hover_color=XANH, command=self.nap_tep_rui_ro).pack(side="left", padx=6)
        ctk.CTkButton(thanh, text="Thêm MST thủ công", width=160, fg_color="#37474F",
                      hover_color="#263238", command=self.them_tay).pack(side="left")
        ctk.CTkButton(thanh, text="Xóa dòng chọn", width=120, fg_color="#78909C",
                      hover_color="#546E7A", command=self.xoa_dong_rui_ro).pack(side="right", padx=6)
        ctk.CTkButton(thanh, text="Xóa toàn bộ", width=110, fg_color=DO,
                      hover_color="#8E0000", command=self.xoa_het_rui_ro).pack(side="right")

        cot = [("mst", "Mã số thuế", 140, "center"),
               ("ten", "Tên doanh nghiệp", 360, "w"),
               ("ly_do", "Lý do / ghi chú", 320, "w"),
               ("nguon", "Nguồn", 180, "w"),
               ("ngay", "Ngày", 110, "center"),
               ("trong_kho", "Có trong kho", 110, "center")]
        self.tv_rr = self._bang(cha, cot, "tv_rr")

        self.nhan_rr = ctk.CTkLabel(cha, text="", font=("Segoe UI", 11), text_color=XAM)
        self.nhan_rr.pack(anchor="w", padx=12, pady=(0, 10))
        self.ve_bang_rui_ro()

    def ve_bang_rui_ro(self):
        if self.tv_rr is None:
            return
        for k in self.tv_rr.get_children(""):
            self.tv_rr.delete(k)
        trong_kho = {n["mst"] for n in self.ds_ncc}
        trung = 0
        for d in self.ds_rui_ro:
            co = d["mst"] in trong_kho
            if co:
                trung += 1
            self.tv_rr.insert("", "end", iid=d["mst"], tags=(RUI_RO if co else KHONG_RO,),
                              values=(d["mst"], d.get("ten", ""), d.get("ly_do", ""),
                                      d.get("nguon", ""), d.get("ngay", ""),
                                      "CÓ" if co else "—"))
        self.nhan_rr.configure(
            text="Danh sách có %d mã số thuế; trong đó %d mã đang là nhà cung cấp trong kho hóa đơn của bạn."
                 % (len(self.ds_rui_ro), trung))

    def _luu_rui_ro(self):
        self.map_rui_ro = {d["mst"]: d for d in self.ds_rui_ro}
        ghi_ds_rui_ro(self.ds_rui_ro)
        self.ve_bang_rui_ro()
        self.ve_bang()

    def nap_tep_rui_ro(self):
        dd = filedialog.askopenfilename(
            title="Chọn tệp danh sách rủi ro",
            filetypes=[("Excel / CSV", "*.xlsx *.xls *.csv"), ("Tất cả", "*.*")])
        if not dd:
            return
        if mst is not None and hasattr(mst, "nap_ds_rui_ro"):
            try:
                mst.nap_ds_rui_ro(dd)
            except Exception:
                pass
        try:
            them = self._doc_tep_rui_ro(dd)
        except Exception:
            messagebox.showerror("Lỗi đọc tệp", traceback.format_exc())
            return
        cu = {d["mst"] for d in self.ds_rui_ro}
        moi = [d for d in them if d["mst"] not in cu]
        self.ds_rui_ro.extend(moi)
        self._luu_rui_ro()
        messagebox.showinfo("Đã nạp", "Đọc được %d dòng, thêm mới %d mã số thuế, bỏ qua %d mã đã có."
                            % (len(them), len(moi), len(them) - len(moi)))

    def _doc_tep_rui_ro(self, duong_dan):
        ds = []
        ngay = dt.datetime.now().strftime("%d/%m/%Y")
        ten_tep = os.path.basename(duong_dan)
        if duong_dan.lower().endswith(".csv"):
            with open(duong_dan, "r", encoding="utf-8-sig", errors="ignore") as f:
                mau = f.read(4096)
                f.seek(0)
                tach = ";" if mau.count(";") > mau.count(",") else ","
                for hang in csv.reader(f, delimiter=tach):
                    d = self._hang_rui_ro(hang, ten_tep, ngay)
                    if d:
                        ds.append(d)
            return ds
        from openpyxl import load_workbook
        wb = load_workbook(duong_dan, read_only=True, data_only=True)
        for ws in wb.worksheets:
            for hang in ws.iter_rows(values_only=True):
                d = self._hang_rui_ro(list(hang), ten_tep, ngay)
                if d:
                    ds.append(d)
        wb.close()
        return ds

    @staticmethod
    def _hang_rui_ro(hang, nguon, ngay):
        o = [("" if v is None else str(v)).strip() for v in hang]
        ma_so = ""
        vt = -1
        for i, v in enumerate(o):
            so = re.sub(r"\D", "", v)
            if len(so) in (10, 13, 14) and so == v.replace("-", "").replace(" ", ""):
                ma_so, vt = so[:10] if len(so) == 10 else so, i
                break
            if len(so) in (10, 12, 13, 14) and len(v) <= 16 and so:
                ma_so, vt = so, i
                break
        if not ma_so or len(ma_so) < 10:
            return None
        con_lai = [v for i, v in enumerate(o) if i != vt and v]
        ten = con_lai[0] if con_lai else ""
        ly_do = " | ".join(con_lai[1:])[:300]
        return {"mst": ma_so, "ten": ten, "ly_do": ly_do, "nguon": nguon, "ngay": ngay}

    def them_tay(self):
        cs = ctk.CTkToplevel(self)
        cs.title("Thêm mã số thuế rủi ro")
        cs.geometry("520x300")
        cs.transient(self.winfo_toplevel())
        cs.grab_set()
        ctk.CTkLabel(cs, text="Mã số thuế (mỗi dòng một mã):",
                     font=("Segoe UI", 12, "bold")).pack(anchor="w", padx=14, pady=(14, 4))
        tb = ctk.CTkTextbox(cs, height=140, font=("Consolas", 11))
        tb.pack(fill="both", expand=True, padx=14)
        ctk.CTkLabel(cs, text="Lý do / ghi chú:", font=("Segoe UI", 12)).pack(anchor="w", padx=14, pady=(8, 2))
        o_ly_do = ctk.CTkEntry(cs, width=480)
        o_ly_do.pack(padx=14)

        def luu():
            ly_do = o_ly_do.get().strip() or "Nhập thủ công"
            ngay = dt.datetime.now().strftime("%d/%m/%Y")
            cu = {d["mst"] for d in self.ds_rui_ro}
            n = 0
            for dong in tb.get("1.0", "end").splitlines():
                ma_so = re.sub(r"\D", "", dong)
                if len(ma_so) >= 10 and ma_so not in cu:
                    self.ds_rui_ro.append({"mst": ma_so, "ten": "", "ly_do": ly_do,
                                           "nguon": "Nhập tay", "ngay": ngay})
                    cu.add(ma_so)
                    n += 1
            self._luu_rui_ro()
            cs.destroy()
            messagebox.showinfo("Đã thêm", "Đã thêm %d mã số thuế." % n)

        hang = ctk.CTkFrame(cs, fg_color="transparent")
        hang.pack(fill="x", padx=14, pady=12)
        ctk.CTkButton(hang, text="Lưu", width=100, fg_color=XANH,
                      hover_color=XANH_DAM, command=luu).pack(side="right")
        ctk.CTkButton(hang, text="Hủy", width=90, fg_color="#607D8B",
                      hover_color="#455A64", command=cs.destroy).pack(side="right", padx=8)

    def xoa_dong_rui_ro(self):
        chon = list(self.tv_rr.selection())
        if not chon:
            messagebox.showinfo("Chưa chọn", "Hãy chọn dòng cần xóa.")
            return
        if not messagebox.askyesno("Xác nhận", "Xóa %d mã số thuế khỏi danh sách rủi ro?" % len(chon)):
            return
        self.ds_rui_ro = [d for d in self.ds_rui_ro if d["mst"] not in set(chon)]
        self._luu_rui_ro()

    def xoa_het_rui_ro(self):
        if not self.ds_rui_ro:
            return
        if not messagebox.askyesno("Xác nhận", "Xóa toàn bộ %d mã số thuế trong danh sách rủi ro?"
                                   % len(self.ds_rui_ro)):
            return
        self.ds_rui_ro = []
        self._luu_rui_ro()

    # ========================= TAB 3: HUONG DAN & CAU HINH ==================

    def _tab_cau_hinh(self, cha):
        khung = ctk.CTkScrollableFrame(cha, fg_color="transparent")
        khung.pack(fill="both", expand=True, padx=6, pady=6)

        ctk.CTkLabel(khung, text="CẤU HÌNH TRA CỨU",
                     font=("Segoe UI", 14, "bold"), text_color=XANH_DAM).pack(anchor="w", padx=8, pady=(6, 8))

        self.bien_bat = tk.BooleanVar(value=bool(self.cau_hinh.get("bat", False)))
        ctk.CTkCheckBox(khung, text="Cho phép tra cứu trực tuyến (mặc định tắt — tắt thì ứng dụng không gửi gì ra ngoài)",
                        variable=self.bien_bat, font=("Segoe UI", 12)).pack(anchor="w", padx=8, pady=4)

        def o_nhap(nhan, gia_tri, rong=520):
            h = ctk.CTkFrame(khung, fg_color="transparent")
            h.pack(fill="x", padx=8, pady=4)
            ctk.CTkLabel(h, text=nhan, width=230, anchor="w", font=("Segoe UI", 12)).pack(side="left")
            e = ctk.CTkEntry(h, width=rong)
            e.pack(side="left")
            e.insert(0, str(gia_tri))
            return e

        self.o_url = o_nhap("Mẫu địa chỉ tra cứu:", self.cau_hinh.get("url_mau", URL_MAC_DINH))
        self.o_ngay = o_nhap("Số ngày dùng lại kết quả cũ:", self.cau_hinh.get("so_ngay_cache", 7), 120)
        self.o_tre = o_nhap("Giãn cách giữa 2 lần gọi (giây):", self.cau_hinh.get("do_tre", 0.6), 120)
        self.o_cho = o_nhap("Thời gian chờ tối đa (giây):", self.cau_hinh.get("thoi_gian_cho", 10), 120)

        hang = ctk.CTkFrame(khung, fg_color="transparent")
        hang.pack(fill="x", padx=8, pady=10)
        ctk.CTkButton(hang, text="LƯU CẤU HÌNH", width=150, fg_color=XANH,
                      hover_color=XANH_DAM, command=self.luu_cau_hinh).pack(side="left")
        ctk.CTkButton(hang, text="Thử một mã số thuế", width=170, fg_color=XANH_DAM,
                      hover_color=XANH, command=self.thu_mot).pack(side="left", padx=6)
        ctk.CTkButton(hang, text="Xóa toàn bộ kết quả đã lưu", width=200, fg_color="#78909C",
                      hover_color="#546E7A", command=self.xoa_cache).pack(side="left")
        ctk.CTkButton(hang, text="Mở thư mục dữ liệu", width=160, fg_color="#455A64",
                      hover_color="#263238", command=self.mo_thu_muc).pack(side="left", padx=6)

        ctk.CTkLabel(khung, text="HƯỚNG DẪN SỬ DỤNG",
                     font=("Segoe UI", 14, "bold"), text_color=XANH_DAM).pack(anchor="w", padx=8, pady=(12, 6))
        tb = ctk.CTkTextbox(khung, height=330, font=("Segoe UI", 12), wrap="word")
        tb.pack(fill="both", expand=True, padx=8, pady=(0, 10))
        tb.insert("1.0", self._van_ban_huong_dan())
        tb.configure(state="disabled")

    @staticmethod
    def _van_ban_huong_dan():
        return (
            "1. NÚT NẠP DANH SÁCH NCC TỪ KHO\n"
            "   Gom toàn bộ người bán có trong kho hóa đơn, kèm số hóa đơn và tổng tiền. "
            "Bước này hoàn toàn ngoại tuyến.\n\n"
            "2. TRA CỨU TÌNH TRẠNG MÃ SỐ THUẾ\n"
            "   Chỉ chạy khi bạn tự bật ở ô phía trên. Kết quả lưu lại trong data\\mst_cache.json, "
            "trong số ngày cấu hình sẽ dùng lại, không gọi lại mạng.\n"
            "   Nguồn mặc định là dịch vụ tra cứu công khai theo mã số thuế. Nhà cung cấp dịch vụ này "
            "thông báo dừng từ ngày 01/03/2027; khi đó bạn chỉ cần sửa lại 'Mẫu địa chỉ tra cứu' "
            "sang nguồn thay thế, phần còn lại giữ nguyên.\n"
            "   Mẫu địa chỉ bắt buộc chứa chuỗi {mst}, ứng dụng sẽ thay bằng mã số thuế thật.\n\n"
            "3. DANH SÁCH DOANH NGHIỆP RỦI RO\n"
            "   Ứng dụng không tự tải tệp của cơ quan thuế. Bạn bấm 'Mở trang Cục Thuế để tải tệp', "
            "tải tệp Excel về máy, rồi bấm 'Nạp tệp Excel/CSV'. Mã số thuế trùng với nhà cung cấp "
            "trong kho sẽ được tô đỏ ở tab thứ nhất.\n\n"
            "4. XÁC MINH LẠI BẰNG TAY\n"
            "   Trang chính thức của cơ quan thuế có mã xác nhận (captcha) nên phải tự mở trình duyệt "
            "và nhập tay. Trong cửa sổ chi tiết mỗi nhà cung cấp có sẵn nút mở trang này.\n\n"
            "5. GIÁ TRỊ PHÁP LÝ\n"
            "   Kết quả tra cứu chỉ để tham khảo khi rà soát nội bộ. Căn cứ chính thức vẫn là tệp XML "
            "gốc của hóa đơn và thông tin tra cứu trực tiếp trên hệ thống của cơ quan thuế.\n\n"
            "TAX INVOICE PRO " + PHIEN_BAN + " — Tác giả: " + TAC_GIA
        )

    def luu_cau_hinh(self):
        url = self.o_url.get().strip() or URL_MAC_DINH
        if "{mst}" not in url:
            messagebox.showerror("Sai mẫu địa chỉ", "Mẫu địa chỉ phải chứa chuỗi {mst}.\n\nVí dụ:\n" + URL_MAC_DINH)
            return

        def so(o, md, kieu):
            try:
                return kieu(str(o.get()).strip().replace(",", "."))
            except Exception:
                return md

        self.cau_hinh.update({
            "bat": bool(self.bien_bat.get()),
            "url_mau": url,
            "so_ngay_cache": max(0, so(self.o_ngay, 7, int)),
            "do_tre": max(0.0, so(self.o_tre, 0.6, float)),
            "thoi_gian_cho": max(3.0, so(self.o_cho, 10.0, float)),
        })
        ghi_cau_hinh(self.cau_hinh)
        messagebox.showinfo("Đã lưu", "Đã lưu cấu hình.\nTra cứu trực tuyến: %s."
                            % ("BẬT" if self.cau_hinh["bat"] else "TẮT"))

    def thu_mot(self):
        from tkinter import simpledialog
        ma_so = simpledialog.askstring("Thử tra cứu", "Nhập một mã số thuế:", parent=self)
        if not ma_so:
            return
        ma_so = re.sub(r"\D", "", ma_so)
        if len(ma_so) < 10:
            messagebox.showerror("Sai định dạng", "Mã số thuế phải có 10, 12 hoặc 13 chữ số.")
            return
        if not self.bien_bat.get():
            messagebox.showwarning("Đang tắt", "Hãy bật 'Cho phép tra cứu trực tuyến' và bấm Lưu cấu hình trước.")
            return
        cfg = dict(self.cau_hinh)
        cfg["url_mau"] = self.o_url.get().strip() or URL_MAC_DINH
        kq = tra_cuu_mot(ma_so, cfg)
        self.cache[ma_so] = kq
        ghi_cache(self.cache)
        self.ve_bang()
        messagebox.showinfo("Kết quả",
                            "Mã số thuế: %s\nTên: %s\nĐịa chỉ: %s\nTình trạng: %s\nĐánh giá: %s"
                            % (ma_so, kq.get("ten") or "(trống)", kq.get("dia_chi") or "(trống)",
                               kq.get("tinh_trang") or "(trống)", NHAN_MUC.get(kq.get("muc"), "")))

    def xoa_cache(self):
        if not messagebox.askyesno("Xác nhận", "Xóa toàn bộ kết quả tra cứu đã lưu?"):
            return
        self.cache = {}
        ghi_cache(self.cache)
        self.ve_bang()
        messagebox.showinfo("Đã xóa", "Đã xóa kết quả tra cứu đã lưu.")

    def mo_thu_muc(self):
        try:
            os.startfile(THU_MUC_DL)
        except Exception:
            messagebox.showinfo("Đường dẫn", THU_MUC_DL)


# Ten goi khac de khoi_dong.py / man_hinh_moi.py nhan dien
ManHinh = ManHinhNCC
ManHinhNhaCungCap = ManHinhNCC
LOP_MAN_HINH = {"ncc": ManHinhNCC}


def mo_cua_so():
    """Chay rieng de thu man hinh nay."""
    ctk.set_appearance_mode("light")
    root = ctk.CTk()
    root.title("TAX INVOICE PRO — Nhà cung cấp & Rủi ro")
    root.geometry("1280x800")
    root.configure(fg_color=NEN)
    ManHinhNCC(root, None).pack(fill="both", expand=True)
    root.mainloop()


if __name__ == "__main__":
    if LOI_MST:
        print("[canh bao] khong nap duoc mst_tra_cuu, dung che do du phong:")
        print(LOI_MST)
    mo_cua_so()


TEP_DAY_DU = True  # dau hieu tep da duoc dan day du - dung xoa dong nay
