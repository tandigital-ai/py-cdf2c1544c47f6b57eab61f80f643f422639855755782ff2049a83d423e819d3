# -*- coding: utf-8 -*-
"""
TAX INVOICE PRO - mst_tra_cuu.py (v1.6.0)
Tra cuu tinh trang hoat dong cua nguoi ban theo ma so thue + danh sach rui ro.
Mac dinh TAT tra cuu truc tuyen. Khi bat chi gui di ma so thue (thong tin cong khai).
Tac gia: TDA - 0986509106
"""
import os
import re
import csv
import json
import time
import sqlite3
import datetime
import threading
import unicodedata
import urllib.request
import urllib.error
import webbrowser
from concurrent.futures import ThreadPoolExecutor

PHIEN_BAN = "1.6.0"
THU_MUC = os.path.dirname(os.path.abspath(__file__))
TEP_CAU_HINH = os.path.join(THU_MUC, "data", "mst.json")

TRANG_TRA_CUU_NNT = "https://tracuunnt.gdt.gov.vn/tcnnt/mstdn.jsp"
TRANG_DN_RUI_RO = "https://tracuuhoadon.gdt.gov.vn/dnrrvp.html"

CAU_HINH_MAC_DINH = {
    "bat_tra_cuu_truc_tuyen": False,
    "url_mau": "https://api.vietqr.io/v2/business/{mst}",
    "url_du_phong": "",
    "timeout": 12,
    "so_luong": 3,
    "do_tre_giay": 0.25,
    "han_cache_ngay": 7,
    "mst_don_vi": "",
}

# Muc rui ro
AN_TOAN = "an_toan"
THEO_DOI = "theo_doi"
RUI_RO = "rui_ro"
KHONG_RO = "khong_ro"

NHAN_MUC = {
    AN_TOAN: "Bình thường",
    THEO_DOI: "Cần theo dõi",
    RUI_RO: "RỦI RO CAO",
    KHONG_RO: "Chưa xác định",
}

_khoa_ghi = threading.Lock()


# ----------------------------------------------------------------------
# Tien ich chung
# ----------------------------------------------------------------------
def _kd(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return s.replace("\u0111", "d").replace("\u0110", "D").lower().strip()


def chuan_mst(v):
    s = re.sub(r"[^0-9]", "", str(v or ""))
    return s


def mst_hop_le(v):
    s = chuan_mst(v)
    return len(s) in (10, 12, 13)


def mst_goc(v):
    """MST 13 so (chi nhanh) -> lay 10 so dau de so sanh voi MST don vi."""
    s = chuan_mst(v)
    return s[:10] if len(s) == 13 else s


def tien(v):
    try:
        return "{:,.0f}".format(float(v or 0)).replace(",", ".")
    except Exception:
        return "0"


def so(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if not s:
        return 0.0
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") \
            else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".") if len(s.split(",")[-1]) in (1, 2, 3) and \
            s.count(",") == 1 and len(s.split(",")[-1]) <= 2 else s.replace(",", "")
    else:
        if s.count(".") > 1:
            s = s.replace(".", "")
    try:
        return float(re.sub(r"[^0-9.\-]", "", s) or 0)
    except Exception:
        return 0.0


def ngay_iso(v):
    s = str(v or "").strip()
    if not s:
        return ""
    s = s.replace("T", " ").split(" ")[0]
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})$", s)
    if m:
        return s
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{4})$", s)
    if m:
        return "%s-%02d-%02d" % (m.group(3), int(m.group(2)), int(m.group(1)))
    return ""


def ngay_vn(v):
    s = ngay_iso(v)
    return "%s/%s/%s" % (s[8:10], s[5:7], s[0:4]) if s else str(v or "")


# ----------------------------------------------------------------------
# Cau hinh
# ----------------------------------------------------------------------
def doc_cau_hinh():
    ch = dict(CAU_HINH_MAC_DINH)
    try:
        if os.path.exists(TEP_CAU_HINH):
            with open(TEP_CAU_HINH, "r", encoding="utf-8") as f:
                ch.update(json.load(f) or {})
    except Exception:
        pass
    if not ch.get("mst_don_vi"):
        for ten in ("config.json", os.path.join("data", "config.json")):
            p = os.path.join(THU_MUC, ten)
            try:
                if os.path.exists(p):
                    with open(p, "r", encoding="utf-8") as f:
                        c = json.load(f) or {}
                    for k in ("mst_don_vi", "mst_mac_dinh", "mst"):
                        if c.get(k):
                            ch["mst_don_vi"] = chuan_mst(c[k])
                            break
            except Exception:
                pass
    return ch


def ghi_cau_hinh(ch):
    try:
        os.makedirs(os.path.dirname(TEP_CAU_HINH), exist_ok=True)
        with open(TEP_CAU_HINH, "w", encoding="utf-8") as f:
            json.dump(ch, f, ensure_ascii=False, indent=2)
        return True
    except Exception:
        return False


# ----------------------------------------------------------------------
# CSDL
# ----------------------------------------------------------------------
def duong_dan_db():
    try:
        import kho
        for ten in ("DUONG_DAN_DB", "TEP_DB", "DB_PATH", "DUONG_DAN_CSDL",
                    "TEP_CSDL", "CSDL", "DB"):
            v = getattr(kho, ten, None)
            if isinstance(v, str) and v.lower().endswith(".db"):
                return v
    except Exception:
        pass
    return os.path.join(THU_MUC, "data", "taxinvoice.db")


def ket_noi():
    p = duong_dan_db()
    os.makedirs(os.path.dirname(p), exist_ok=True)
    conn = sqlite3.connect(p, timeout=30)
    conn.row_factory = sqlite3.Row
    return conn


def _bang_co(conn):
    return {r[0] for r in conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table'")}


def _cot_cua(conn, bang):
    try:
        return [r[1] for r in conn.execute("PRAGMA table_info(%s)" % bang)]
    except Exception:
        return []


def ten_bang_hoa_don(conn=None):
    tu_dong = conn is None
    conn = conn or ket_noi()
    try:
        b = _bang_co(conn)
        for t in ("hoa_don", "invoices", "hoadon"):
            if t in b:
                return t
        return ""
    finally:
        if tu_dong:
            conn.close()


def khoi_tao_bang():
    conn = ket_noi()
    try:
        conn.execute("""CREATE TABLE IF NOT EXISTS nha_cung_cap(
            mst TEXT PRIMARY KEY, ten TEXT, ten_quoc_te TEXT, dia_chi TEXT,
            trang_thai TEXT, ma_trang_thai TEXT, muc_rui_ro TEXT,
            nguon TEXT, ngay_tra_cuu TEXT, ghi_chu TEXT, du_lieu TEXT)""")
        conn.execute("""CREATE TABLE IF NOT EXISTS ds_rui_ro(
            mst TEXT PRIMARY KEY, ten TEXT, nguon TEXT, ngay_nhap TEXT,
            ghi_chu TEXT)""")
        bang = ten_bang_hoa_don(conn)
        if bang:
            cot = _cot_cua(conn, bang)
            for ten, kieu in (("tt_khong_tien_mat", "INTEGER DEFAULT 0"),
                              ("ghi_chu_kt", "TEXT")):
                if ten not in cot:
                    try:
                        conn.execute("ALTER TABLE %s ADD COLUMN %s %s"
                                     % (bang, ten, kieu))
                    except Exception:
                        pass
        conn.commit()
    finally:
        conn.close()


# ----------------------------------------------------------------------
# Doc hoa don (tu do ten cot)
# ----------------------------------------------------------------------
BAN_DO = {
    "id": ("id", "ma", "rowid"),
    "khoa": ("khoa", "khoa_trung", "key"),
    "ngay_lap": ("ngay_lap", "ngay", "ngay_hoa_don", "nlap", "ngaylap"),
    "ngay_ky": ("ngay_ky", "thoi_diem_ky", "signing_time", "ngay_ky_so"),
    "so_hd": ("so_hoa_don", "so_hd", "shdon", "so"),
    "ky_hieu": ("ky_hieu_hoa_don", "ky_hieu", "khhdon"),
    "mau_so": ("ky_hieu_mau_so", "mau_so", "khmshdon"),
    "mst_ban": ("mst_nguoi_ban", "mst_ban", "mstnban", "mst_seller"),
    "ten_ban": ("ten_nguoi_ban", "ten_ban", "tennban"),
    "mst_mua": ("mst_nguoi_mua", "mst_mua", "mstnmua"),
    "ten_mua": ("ten_nguoi_mua", "ten_mua", "tennmua"),
    "tien_hang": ("tien_hang", "tgtcthue", "tong_tien_hang"),
    "thue": ("thue_gtgt", "tien_thue", "tgtthue"),
    "tong": ("tong_tien", "tong_thanh_toan", "tgttttbso"),
    "trang_thai": ("trang_thai", "tthai", "tt"),
    "loai_hd": ("loai_hoa_don", "loai_hd", "loai"),
    "xml": ("duong_dan_xml", "tep_xml", "duong_dan", "file_xml"),
    "tt_kdtm": ("tt_khong_tien_mat",),
    "ghi_chu": ("ghi_chu_kt", "ghi_chu"),
}


def doc_hoa_don(tu_ngay=None, den_ngay=None):
    """Tra ve danh sach dict da chuan hoa khoa, khong phu thuoc ten cot goc."""
    conn = ket_noi()
    try:
        bang = ten_bang_hoa_don(conn)
        if not bang:
            return []
        cot = _cot_cua(conn, bang)
        thap = {c.lower(): c for c in cot}
        rows = conn.execute("SELECT * FROM %s" % bang).fetchall()
    except Exception:
        return []
    finally:
        conn.close()

    def lay(d, khoa):
        for ten in BAN_DO[khoa]:
            if ten in thap:
                v = d.get(thap[ten])
                if v not in (None, ""):
                    return v
        return ""

    ds = []
    for r in rows:
        d = dict(r)
        hd = {
            "id": lay(d, "id"),
            "ngay_lap": ngay_iso(lay(d, "ngay_lap")),
            "ngay_ky": ngay_iso(lay(d, "ngay_ky")),
            "so_hd": str(lay(d, "so_hd") or "").strip(),
            "ky_hieu": str(lay(d, "ky_hieu") or "").strip(),
            "mau_so": str(lay(d, "mau_so") or "").strip(),
            "mst_ban": chuan_mst(lay(d, "mst_ban")),
            "ten_ban": str(lay(d, "ten_ban") or "").strip(),
            "mst_mua": chuan_mst(lay(d, "mst_mua")),
            "ten_mua": str(lay(d, "ten_mua") or "").strip(),
            "tien_hang": so(lay(d, "tien_hang")),
            "thue": so(lay(d, "thue")),
            "tong": so(lay(d, "tong")),
            "trang_thai": str(lay(d, "trang_thai") or "").strip(),
            "loai_hd": str(lay(d, "loai_hd") or "").strip(),
            "xml": str(lay(d, "xml") or "").strip(),
            "tt_kdtm": 1 if str(lay(d, "tt_kdtm") or "0") in ("1", "True", "true") else 0,
            "ghi_chu": str(lay(d, "ghi_chu") or "").strip(),
        }
        k = lay(d, "khoa")
        hd["khoa"] = str(k) if k else "%s|%s|%s" % (hd["mst_ban"], hd["ky_hieu"],
                                                    hd["so_hd"])
        if hd["tong"] == 0 and (hd["tien_hang"] or hd["thue"]):
            hd["tong"] = hd["tien_hang"] + hd["thue"]
        if tu_ngay and hd["ngay_lap"] and hd["ngay_lap"] < tu_ngay:
            continue
        if den_ngay and hd["ngay_lap"] and hd["ngay_lap"] > den_ngay:
            continue
        ds.append(hd)
    ds.sort(key=lambda x: (x["ngay_lap"], x["so_hd"]))
    return ds


def dat_chung_tu_thanh_toan(khoa, co=True):
    """Danh dau hoa don da co chung tu thanh toan khong dung tien mat."""
    conn = ket_noi()
    try:
        bang = ten_bang_hoa_don(conn)
        cot = [c.lower() for c in _cot_cua(conn, bang)]
        if "tt_khong_tien_mat" not in cot:
            khoi_tao_bang()
        cot_khoa = None
        for c in ("khoa", "khoa_trung", "id"):
            if c in cot:
                cot_khoa = c
                break
        if not cot_khoa:
            return False
        if cot_khoa == "id":
            return False
        conn.execute("UPDATE %s SET tt_khong_tien_mat=? WHERE %s=?"
                     % (bang, cot_khoa), (1 if co else 0, khoa))
        conn.commit()
        return True
    except Exception:
        return False
    finally:
        conn.close()


# ----------------------------------------------------------------------
# Phan loai trang thai MST
# ----------------------------------------------------------------------
def phan_loai_trang_thai(van_ban):
    t = _kd(van_ban)
    if not t:
        return ("", "Chưa tra cứu", KHONG_RO)
    if "khong hoat dong tai dia chi" in t or "bo dia chi" in t:
        return ("06", "06 - Không hoạt động tại địa chỉ đã đăng ký", RUI_RO)
    if "pha san" in t:
        return ("07", "07 - Chờ làm thủ tục phá sản", RUI_RO)
    if "tam ngung" in t or "tam nghi" in t:
        return ("05", "05 - Tạm ngừng hoạt động, kinh doanh", THEO_DOI)
    if "ngung hoat dong" in t and ("chua hoan thanh" in t or "chua dong" in t):
        return ("03", "03 - Ngừng hoạt động nhưng chưa đóng MST", RUI_RO)
    if "ngung hoat dong" in t or "cham dut hieu luc" in t or "da dong ma so thue" in t:
        return ("01", "01 - Ngừng hoạt động, đã đóng MST", RUI_RO)
    if "chuyen co quan thue" in t or "chuyen dia diem" in t:
        return ("02", "02 - Đã chuyển cơ quan thuế quản lý", THEO_DOI)
    if "dang hoat dong" in t or "hoat dong" in t:
        return ("00", "00 - NNT đang hoạt động", AN_TOAN)
    return ("", str(van_ban).strip(), KHONG_RO)


# ----------------------------------------------------------------------
# Goi API
# ----------------------------------------------------------------------
def _http_json(url, timeout):
    req = urllib.request.Request(url, headers={
        "User-Agent": "TaxInvoicePro/%s (ke toan noi bo)" % PHIEN_BAN,
        "Accept": "application/json, text/plain, */*",
    })
    with urllib.request.urlopen(req, timeout=timeout) as r:
        raw = r.read().decode("utf-8", "replace")
    return json.loads(raw)


def _tim_khoa(d, *ten):
    """Tim de quy mot khoa trong json tra ve."""
    muon = {_kd(t) for t in ten}
    nguon = [d]
    while nguon:
        x = nguon.pop(0)
        if isinstance(x, dict):
            for k, v in x.items():
                if _kd(k) in muon and isinstance(v, (str, int, float)) and str(v).strip():
                    return str(v).strip()
                if isinstance(v, (dict, list)):
                    nguon.append(v)
        elif isinstance(x, list):
            nguon.extend([i for i in x if isinstance(i, (dict, list))])
    return ""


def _rec_trong(mst):
    return {"mst": mst, "ten": "", "ten_quoc_te": "", "dia_chi": "",
            "trang_thai": "", "ma_trang_thai": "", "muc_rui_ro": KHONG_RO,
            "nguon": "", "ngay_tra_cuu": "", "ghi_chu": "", "du_lieu": ""}


def tra_cuu_mot(mst, ch=None):
    """Goi API cho 1 MST. Luon tra ve dict, khong nem loi ra ngoai."""
    ch = ch or doc_cau_hinh()
    mst = chuan_mst(mst)
    rec = _rec_trong(mst)
    if not mst:
        rec["ghi_chu"] = "MST rỗng"
        return rec
    urls = [u for u in (ch.get("url_mau"), ch.get("url_du_phong")) if u]
    loi_cuoi = ""
    for u in urls:
        try:
            js = _http_json(u.format(mst=mst), int(ch.get("timeout", 12)))
        except urllib.error.HTTPError as e:
            loi_cuoi = "HTTP %s" % e.code
            if e.code == 404:
                rec["trang_thai"] = "Không tìm thấy MST trên nguồn tra cứu"
                rec["muc_rui_ro"] = KHONG_RO
                rec["nguon"] = u.split("/")[2] if "//" in u else u
                rec["ngay_tra_cuu"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
                return rec
            continue
        except Exception as e:
            loi_cuoi = str(e)
            continue
        ten = _tim_khoa(js, "name", "ten", "tenNNT", "companyName")
        tt = _tim_khoa(js, "status", "tinhTrang", "trangThai", "tinhTrangHoatDong")
        rec["ten"] = ten
        rec["ten_quoc_te"] = _tim_khoa(js, "internationalName")
        rec["dia_chi"] = _tim_khoa(js, "address", "diaChi")
        rec["trang_thai"] = tt or ("Đang hoạt động" if ten else "")
        ma, nhan, muc = phan_loai_trang_thai(rec["trang_thai"])
        rec["ma_trang_thai"] = ma
        rec["trang_thai"] = nhan
        rec["muc_rui_ro"] = muc
        rec["nguon"] = u.split("/")[2] if "//" in u else u
        rec["ngay_tra_cuu"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
        try:
            rec["du_lieu"] = json.dumps(js, ensure_ascii=False)[:4000]
        except Exception:
            pass
        if not ten:
            rec["ghi_chu"] = "Nguồn không trả về tên - cần tra cứu lại trên tracuunnt"
        return rec
    rec["ghi_chu"] = "Không gọi được nguồn tra cứu (%s)" % (loi_cuoi or "không rõ")
    rec["ngay_tra_cuu"] = ""
    return rec


def _tra_mot_an_toan(mst, ch, dung=None):
    if dung is not None and dung.is_set():
        r = _rec_trong(chuan_mst(mst))
        r["ghi_chu"] = "Đã dừng theo yêu cầu"
        return r
    try:
        r = tra_cuu_mot(mst, ch)
    except Exception as e:
        r = _rec_trong(chuan_mst(mst))
        r["ghi_chu"] = "Lỗi: %s" % e
    try:
        time.sleep(float(ch.get("do_tre_giay", 0.25)))
    except Exception:
        pass
    return r


# ----------------------------------------------------------------------
# Cache + danh sach rui ro
# ----------------------------------------------------------------------
def _luu_cache(conn, rec):
    with _khoa_ghi:
        conn.execute("""INSERT OR REPLACE INTO nha_cung_cap
            (mst,ten,ten_quoc_te,dia_chi,trang_thai,ma_trang_thai,muc_rui_ro,
             nguon,ngay_tra_cuu,ghi_chu,du_lieu)
            VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                     (rec.get("mst"), rec.get("ten"), rec.get("ten_quoc_te"),
                      rec.get("dia_chi"), rec.get("trang_thai"),
                      rec.get("ma_trang_thai"), rec.get("muc_rui_ro"),
                      rec.get("nguon"), rec.get("ngay_tra_cuu"),
                      rec.get("ghi_chu"), rec.get("du_lieu")))


def doc_cache_tat_ca():
    khoi_tao_bang()
    conn = ket_noi()
    try:
        kq = {}
        for r in conn.execute("SELECT * FROM nha_cung_cap"):
            kq[r["mst"]] = dict(r)
        return ap_dung_rui_ro(kq)
    except Exception:
        return {}
    finally:
        conn.close()


def doc_ds_rui_ro():
    khoi_tao_bang()
    conn = ket_noi()
    try:
        return {r["mst"]: dict(r) for r in conn.execute("SELECT * FROM ds_rui_ro")}
    except Exception:
        return {}
    finally:
        conn.close()


def ap_dung_rui_ro(kq):
    rr = doc_ds_rui_ro()
    if not rr:
        return kq
    for m, rec in kq.items():
        if m in rr:
            rec["muc_rui_ro"] = RUI_RO
            g = rec.get("ghi_chu") or ""
            them = "Có tên trong danh sách doanh nghiệp rủi ro/vi phạm đã nạp"
            if them not in g:
                rec["ghi_chu"] = (g + " | " if g else "") + them
    return kq


def _het_han(ngay, so_ngay):
    if not ngay:
        return True
    try:
        d = datetime.datetime.strptime(str(ngay)[:10], "%Y-%m-%d")
        return (datetime.datetime.now() - d).days >= int(so_ngay)
    except Exception:
        return True


def tra_cuu_nhieu(ds_mst, lam_moi=False, tien_trinh=None, dung=None, ch=None):
    """Tra cuu hang loat. tien_trinh(i, tong, rec) duoc goi sau moi MST."""
    ch = ch or doc_cau_hinh()
    khoi_tao_bang()
    ds = [m for m in dict.fromkeys([chuan_mst(x) for x in ds_mst]) if m]
    cache = doc_cache_tat_ca()
    han = int(ch.get("han_cache_ngay", 7))
    kq, can_tra = {}, []
    for m in ds:
        c = cache.get(m)
        if (not lam_moi) and c and c.get("ngay_tra_cuu") and \
                not _het_han(c.get("ngay_tra_cuu"), han):
            kq[m] = c
        else:
            can_tra.append(m)
    tong, i = len(ds), 0
    for m in list(kq):
        i += 1
        if tien_trinh:
            tien_trinh(i, tong, kq[m])
    if not ch.get("bat_tra_cuu_truc_tuyen"):
        for m in can_tra:
            i += 1
            r = cache.get(m) or _rec_trong(m)
            r["ghi_chu"] = "Chưa bật tra cứu trực tuyến (xem tab Hướng dẫn & Cấu hình)"
            kq[m] = r
            if tien_trinh:
                tien_trinh(i, tong, r)
        return ap_dung_rui_ro(kq)
    conn = ket_noi()
    try:
        n = max(1, min(6, int(ch.get("so_luong", 3))))
        with ThreadPoolExecutor(max_workers=n) as ex:
            for r in ex.map(lambda m: _tra_mot_an_toan(m, ch, dung), can_tra):
                i += 1
                if r.get("ngay_tra_cuu"):
                    try:
                        _luu_cache(conn, r)
                    except Exception:
                        pass
                kq[r["mst"]] = r
                if tien_trinh:
                    tien_trinh(i, tong, r)
        conn.commit()
    finally:
        conn.close()
    return ap_dung_rui_ro(kq)


# ----------------------------------------------------------------------
# Nap danh sach doanh nghiep rui ro tu tep cua Cuc Thue
# ----------------------------------------------------------------------
def nap_ds_rui_ro(duong_dan, nguon=""):
    khoi_tao_bang()
    nguon = nguon or os.path.basename(duong_dan)
    dong = []
    ext = os.path.splitext(duong_dan)[1].lower()
    try:
        if ext in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            wb = load_workbook(duong_dan, read_only=True, data_only=True)
            for ws in wb.worksheets:
                for r in ws.iter_rows(values_only=True):
                    dong.append(["" if c is None else str(c) for c in r])
            wb.close()
        else:
            for ma in ("utf-8-sig", "utf-8", "cp1258", "latin-1"):
                try:
                    with open(duong_dan, "r", encoding=ma, newline="") as f:
                        mau = f.read(4096)
                        f.seek(0)
                        dau = ";" if mau.count(";") > mau.count(",") else ","
                        dong = [list(r) for r in csv.reader(f, delimiter=dau)]
                    break
                except Exception:
                    continue
    except Exception as e:
        return (0, "Không đọc được tệp: %s" % e)
    if not dong:
        return (0, "Tệp rỗng hoặc không đọc được.")

    cot_mst, cot_ten = -1, -1
    for r in dong[:40]:
        for j, c in enumerate(r):
            t = _kd(c)
            if cot_mst < 0 and ("ma so thue" in t or t == "mst" or "masothue" in t):
                cot_mst = j
            if cot_ten < 0 and ("ten nguoi nop thue" in t or "ten doanh nghiep" in t
                                or t == "ten" or "ten nnt" in t or "ten dn" in t):
                cot_ten = j
        if cot_mst >= 0:
            break
    nay = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    conn = ket_noi()
    n = 0
    try:
        for r in dong:
            mst = ""
            if cot_mst >= 0 and cot_mst < len(r):
                mst = chuan_mst(r[cot_mst])
            if not mst_hop_le(mst):
                mst = ""
                for c in r:
                    x = chuan_mst(c)
                    if mst_hop_le(x):
                        mst = x
                        break
            if not mst_hop_le(mst):
                continue
            ten = ""
            if 0 <= cot_ten < len(r):
                ten = str(r[cot_ten]).strip()
            if not ten:
                for c in r:
                    s = str(c).strip()
                    if len(s) > 8 and not s.isdigit():
                        ten = s
                        break
            conn.execute("""INSERT OR REPLACE INTO ds_rui_ro
                (mst,ten,nguon,ngay_nhap,ghi_chu) VALUES(?,?,?,?,?)""",
                         (mst, ten[:250], nguon, nay, ""))
            n += 1
        conn.commit()
    finally:
        conn.close()
    return (n, "Đã nạp %d mã số thuế từ %s." % (n, nguon))


def xoa_ds_rui_ro():
    conn = ket_noi()
    try:
        conn.execute("DELETE FROM ds_rui_ro")
        conn.commit()
        return True
    finally:
        conn.close()


# ----------------------------------------------------------------------
# Tong hop nha cung cap tu kho hoa don
# ----------------------------------------------------------------------
def tong_hop_ncc(ds_hd, mst_don_vi=""):
    dv = mst_goc(mst_don_vi)
    gom = {}
    for hd in ds_hd:
        m = hd.get("mst_ban") or ""
        if not m:
            continue
        if dv and mst_goc(m) == dv:
            continue  # hoa don ban ra cua chinh don vi
        g = gom.setdefault(m, {"mst": m, "ten_hd": hd.get("ten_ban", ""),
                               "so_hd": 0, "tien_hang": 0.0, "thue": 0.0,
                               "tong": 0.0, "dau": "", "cuoi": ""})
        g["so_hd"] += 1
        g["tien_hang"] += hd.get("tien_hang", 0)
        g["thue"] += hd.get("thue", 0)
        g["tong"] += hd.get("tong", 0)
        n = hd.get("ngay_lap", "")
        if n:
            g["dau"] = n if not g["dau"] else min(g["dau"], n)
            g["cuoi"] = n if not g["cuoi"] else max(g["cuoi"], n)
    return sorted(gom.values(), key=lambda x: -x["tong"])


def xuat_excel_ncc(dong, duong_dan):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "Tinh trang NCC"
    tieu = ["MST", "Tên trên hóa đơn", "Tên theo cơ quan thuế", "Địa chỉ",
            "Tình trạng MST", "Mức rủi ro", "Số hóa đơn", "Tiền hàng",
            "Thuế GTGT", "Tổng thanh toán", "Giao dịch đầu", "Giao dịch cuối",
            "Nguồn tra cứu", "Thời điểm tra cứu", "Ghi chú"]
    ws.append(tieu)
    vien = Border(*[Side(style="thin", color="B0BEC5")] * 4)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1565C0")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = vien
    for d in dong:
        ws.append([d.get("mst", ""), d.get("ten_hd", ""), d.get("ten", ""),
                   d.get("dia_chi", ""), d.get("trang_thai", ""),
                   NHAN_MUC.get(d.get("muc_rui_ro", KHONG_RO), ""),
                   d.get("so_hd", 0), d.get("tien_hang", 0), d.get("thue", 0),
                   d.get("tong", 0), ngay_vn(d.get("dau", "")),
                   ngay_vn(d.get("cuoi", "")), d.get("nguon", ""),
                   d.get("ngay_tra_cuu", ""), d.get("ghi_chu", "")])
    for i, w in enumerate([16, 40, 40, 44, 34, 14, 10, 16, 14, 16, 13, 13, 18, 18, 46], 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.border = vien
            if c.column in (8, 9, 10):
                c.number_format = "#,##0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws2 = wb.create_sheet("Ghi chu")
    for d in ["TAX INVOICE PRO %s - Tác giả: TDA - 0986509106" % PHIEN_BAN,
              "",
              "Kết quả tra cứu chỉ mang tính tham khảo nội bộ.",
              "Chứng cứ pháp lý là bản tra cứu tại:",
              TRANG_TRA_CUU_NNT,
              TRANG_DN_RUI_RO,
              "",
              "Trạng thái MST cần lưu ý: 01, 03, 06, 07 (ngừng hoạt động, bỏ địa chỉ,",
              "phá sản) và 05 (tạm ngừng). Phải đối chiếu ngày lập hóa đơn với thời",
              "điểm MST đổi trạng thái trước khi kết luận hóa đơn không hợp pháp."]:
        ws2.append([d])
    ws2.column_dimensions["A"].width = 90
    wb.save(duong_dan)
    return duong_dan


def mo_trang_nnt():
    webbrowser.open(TRANG_TRA_CUU_NNT)


def mo_trang_rui_ro():
    webbrowser.open(TRANG_DN_RUI_RO)


if __name__ == "__main__":
    khoi_tao_bang()
    print("CSDL   :", duong_dan_db())
    hd = doc_hoa_don()
    print("Hoa don:", len(hd))
    ncc = tong_hop_ncc(hd, doc_cau_hinh().get("mst_don_vi", ""))
    print("NCC    :", len(ncc))
    for d in ncc[:10]:
        print("  %-14s %-45s %3d HD  %15s" %
              (d["mst"], d["ten_hd"][:45], d["so_hd"], tien(d["tong"])))
