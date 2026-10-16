# -*- coding: utf-8 -*-
"""
man_hinh_moi.py - TAX INVOICE PRO v1.5.0
Ba man hinh gan vao app.py: Trang chu moi (bieu do + canh bao),
Doi chieu & Kiem tra (4 tab), Bao cao & Ke khai (2 tab).
Tac gia: TDA - 0986509106
"""
import os
import datetime
import threading
import traceback
import webbrowser
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

try:
    import va_ttk
except Exception:
    pass

import customtkinter as ctk

import doi_chieu as dc
import bao_cao as bc
import bao_cao_chi_tiet as bct
import theo_doi
import hop_thu

XANH, XANH_DAM, NEN = "#1565C0", "#0D47A1", "#F4F7FB"
LUC, CAM, DO, XAM, TIM = "#2E7D32", "#EF6C00", "#C62828", "#607D8B", "#6A1B9A"
THU_MUC = os.path.dirname(os.path.abspath(__file__))
MAU_COT = ["#1565C0", "#1E88E5", "#42A5F5", "#5E9BE6"]


def dat_style():
    s = ttk.Style()
    try:
        s.theme_use("clam")
    except Exception:
        pass
    s.configure("M.Treeview", rowheight=25, font=("Segoe UI", 9),
                background="white", fieldbackground="white")
    s.configure("M.Treeview.Heading", font=("Segoe UI", 9, "bold"),
                background=XANH, foreground="white", relief="flat")
    s.map("M.Treeview.Heading", background=[("active", XANH_DAM)])


def tao_bang(cha, cot, cao=12):
    kh = ctk.CTkFrame(cha, fg_color="white")
    tv = ttk.Treeview(kh, columns=[c[0] for c in cot], show="headings",
                      height=cao, style="M.Treeview")
    for ma, nhan, w in cot:
        tv.heading(ma, text=nhan)
        tv.column(ma, width=w, anchor="w", stretch=(w >= 220))
    d = ttk.Scrollbar(kh, orient="vertical", command=tv.yview)
    n = ttk.Scrollbar(kh, orient="horizontal", command=tv.xview)
    tv.configure(yscrollcommand=d.set, xscrollcommand=n.set)
    tv.grid(row=0, column=0, sticky="nsew")
    d.grid(row=0, column=1, sticky="ns")
    n.grid(row=1, column=0, sticky="ew")
    kh.grid_rowconfigure(0, weight=1)
    kh.grid_columnconfigure(0, weight=1)
    for t, m in (("nang", "#FFEBEE"), ("vua", "#FFF8E1"),
                 ("nhe", "#F1F8E9"), ("ban", "#E8F5E9")):
        tv.tag_configure(t, background=m)
    return kh, tv


def nut(cha, text, lenh, mau=XANH, w=150, h=32):
    return ctk.CTkButton(cha, text=text, command=lenh, width=w, height=h,
                         fg_color=mau, hover_color=XANH_DAM,
                         font=ctk.CTkFont("Segoe UI", 12, "bold"))


def tien_ngan(v):
    v = float(v or 0)
    if abs(v) >= 1e9:
        return "%.2f tỷ" % (v / 1e9)
    if abs(v) >= 1e6:
        return "%.1f tr" % (v / 1e6)
    if abs(v) >= 1e3:
        return "%.0f ng" % (v / 1e3)
    return "%.0f" % v


class ManHinhNen(ctk.CTkFrame):
    """Khung chung: tieu de + dong trang thai."""

    def __init__(self, master, ung_dung=None, tieu_de="", phu=""):
        super().__init__(master, fg_color=NEN, corner_radius=0)
        self.ung_dung = ung_dung
        dat_style()
        dau = ctk.CTkFrame(self, fg_color="transparent")
        dau.pack(fill="x", padx=18, pady=(14, 2))
        ctk.CTkLabel(dau, text=tieu_de,
                     font=ctk.CTkFont("Segoe UI", 20, "bold"),
                     text_color=XANH_DAM).pack(anchor="w")
        if phu:
            ctk.CTkLabel(dau, text=phu, text_color="#5A6B7B",
                         font=ctk.CTkFont("Segoe UI", 11)).pack(anchor="w")
        self.lb_tt = ctk.CTkLabel(self, text="", text_color="#44546A",
                                  font=ctk.CTkFont("Segoe UI", 10))
        self.lb_tt.pack(side="bottom", anchor="w", padx=18, pady=(0, 6))

    def bao(self, t):
        try:
            self.lb_tt.configure(text=t)
            self.update_idletasks()
        except Exception:
            pass


# ======================================================== TRANG CHU MOI
class TrangChuMoi(ManHinhNen):
    def __init__(self, master, ung_dung=None):
        super().__init__(master, ung_dung, "Trang chủ",
                         "Tổng quan kho hóa đơn điện tử đang lưu trên máy này")
        self.ds, self.canh_bao = [], []
        cuon = ctk.CTkScrollableFrame(self, fg_color=NEN)
        cuon.pack(fill="both", expand=True, padx=12, pady=6)
        self.cuon = cuon

        self.the = {}
        hang = ctk.CTkFrame(cuon, fg_color="transparent")
        hang.pack(fill="x", pady=(4, 10))
        for i, (ma, nhan, mau) in enumerate(
                [("so", "Tổng số hóa đơn", XANH_DAM),
                 ("hang_tien", "Tổng tiền hàng", LUC),
                 ("thue", "Tổng thuế GTGT", CAM),
                 ("tong", "Tổng thanh toán", XANH)]):
            k = ctk.CTkFrame(hang, fg_color="white", corner_radius=8)
            k.grid(row=0, column=i, sticky="nsew", padx=(0 if i == 0 else 10, 0))
            hang.grid_columnconfigure(i, weight=1)
            ctk.CTkLabel(k, text=nhan, text_color="#5A6B7B",
                         font=ctk.CTkFont("Segoe UI", 11)).pack(anchor="w",
                                                                padx=14,
                                                                pady=(12, 0))
            l = ctk.CTkLabel(k, text="0", text_color=mau,
                             font=ctk.CTkFont("Segoe UI", 22, "bold"))
            l.pack(anchor="w", padx=14, pady=(0, 12))
            self.the[ma] = l

        kb = ctk.CTkFrame(cuon, fg_color="white", corner_radius=8)
        kb.pack(fill="x", pady=(0, 10))
        ctk.CTkLabel(kb, text="Tổng thanh toán theo tháng",
                     font=ctk.CTkFont("Segoe UI", 13, "bold"),
                     text_color=XANH_DAM).pack(anchor="w", padx=14, pady=(10, 0))
        self.canvas = tk.Canvas(kb, height=250, bg="white", highlightthickness=0)
        self.canvas.pack(fill="x", padx=10, pady=(4, 12))
        self.canvas.bind("<Configure>", lambda e: self.ve_bieu_do())

        kc = ctk.CTkFrame(cuon, fg_color="white", corner_radius=8)
        kc.pack(fill="both", expand=True)
        t = ctk.CTkFrame(kc, fg_color="transparent")
        t.pack(fill="x", padx=14, pady=(10, 4))
        ctk.CTkLabel(t, text="Cảnh báo cần xử lý",
                     font=ctk.CTkFont("Segoe UI", 13, "bold"),
                     text_color=XANH_DAM).pack(side="left")
        self.lb_cb = ctk.CTkLabel(t, text="", text_color=DO,
                                  font=ctk.CTkFont("Segoe UI", 12, "bold"))
        self.lb_cb.pack(side="left", padx=14)
        nut(t, "Kiểm tra lại", self.kiem_tra, XAM, 120).pack(side="right")
        nut(t, "Mở màn hình đối chiếu",
            lambda: self.di_toi("doi_chieu"), XANH, 190).pack(side="right",
                                                              padx=8)
        kh, self.tv = tao_bang(kc, [("muc", "Mức", 60), ("loai", "Loại", 130),
                                    ("hd", "Hóa đơn", 220),
                                    ("mo_ta", "Nội dung", 640)], 9)
        kh.pack(fill="both", expand=True, padx=14, pady=(0, 12))

        t2 = ctk.CTkFrame(cuon, fg_color="transparent")
        t2.pack(fill="x", pady=10)
        nut(t2, "Tra cứu hóa đơn", lambda: self.di_toi("tra_cuu"),
            XANH, 170).pack(side="left")
        nut(t2, "Nạp hóa đơn từ máy", self.nap, LUC, 190).pack(side="left",
                                                               padx=8)
        nut(t2, "Báo cáo & Kê khai", lambda: self.di_toi("bao_cao"),
            TIM, 180).pack(side="left")
        nut(t2, "Làm mới", self.lam_moi, XAM, 110).pack(side="left", padx=8)

        self.after(200, self.lam_moi)

    def di_toi(self, ma):
        if self.ung_dung and hasattr(self.ung_dung, "hien_man_hinh"):
            self.ung_dung.hien_man_hinh(ma)

    def nap(self):
        for ten in ("mo_hop_nap", "mo_nap", "tai_hoa_don"):
            h = getattr(self.ung_dung, ten, None)
            if callable(h):
                try:
                    return h()
                except Exception:
                    pass
        self.di_toi("tra_cuu")

    def lam_moi(self):
        try:
            self.ds = dc.doc_hoa_don()
        except Exception as e:
            self.bao("Không đọc được kho: %s" % e)
            return
        self.the["so"].configure(text="%d" % len(self.ds))
        self.the["hang_tien"].configure(
            text=dc.fmt_tien(sum(h["tien_hang"] for h in self.ds)))
        self.the["thue"].configure(
            text=dc.fmt_tien(sum(h["thue_gtgt"] for h in self.ds)))
        self.the["tong"].configure(
            text=dc.fmt_tien(sum(h["tong_tien"] for h in self.ds)))
        self.ve_bieu_do()
        self.bao("Cập nhật lúc %s — %d hóa đơn."
                 % (datetime.datetime.now().strftime("%H:%M:%S"), len(self.ds)))
        self.after(150, self.kiem_tra)

    def _theo_thang(self):
        g = {}
        for h in self.ds:
            k = (h["ngay_lap"] or "")[:7]
            if len(k) != 7:
                continue
            d = g.setdefault(k, [0, 0])
            d[0] += h["tong_tien"]
            d[1] += 1
        return sorted(g.items())[-12:]

    def ve_bieu_do(self):
        c = self.canvas
        c.delete("all")
        w = c.winfo_width() or 900
        h = 250
        dl = self._theo_thang()
        if not dl:
            c.create_text(w / 2, h / 2, text="Chưa có dữ liệu để vẽ biểu đồ",
                          fill="#8A97A5", font=("Segoe UI", 12))
            return
        trai, phai, tren, duoi = 70, 24, 24, 42
        cao = h - tren - duoi
        rong = max(1, w - trai - phai)
        lon = max(v[0] for _, v in dl) or 1
        for i in range(5):
            y = tren + cao - cao * i / 4.0
            c.create_line(trai, y, w - phai, y, fill="#E3E9F0")
            c.create_text(trai - 8, y, anchor="e", fill="#8A97A5",
                          font=("Segoe UI", 8),
                          text=tien_ngan(lon * i / 4.0))
        n = len(dl)
        bw = rong / float(n)
        for i, (ky, (tien, sl)) in enumerate(dl):
            x0 = trai + i * bw + bw * 0.22
            x1 = trai + (i + 1) * bw - bw * 0.22
            ch = cao * (tien / float(lon))
            y1 = tren + cao
            y0 = y1 - max(2, ch)
            c.create_rectangle(x0, y0, x1, y1, outline="",
                               fill=MAU_COT[i % len(MAU_COT)])
            c.create_text((x0 + x1) / 2, y0 - 9, text=tien_ngan(tien),
                          fill=XANH_DAM, font=("Segoe UI", 8, "bold"))
            c.create_text((x0 + x1) / 2, y1 + 12,
                          text="%s/%s" % (ky[5:7], ky[2:4]),
                          fill="#5A6B7B", font=("Segoe UI", 8))
            c.create_text((x0 + x1) / 2, y1 + 26, text="%d HĐ" % sl,
                          fill="#8A97A5", font=("Segoe UI", 8))
        c.create_line(trai, tren + cao, w - phai, tren + cao, fill="#B8C6D6")

    def kiem_tra(self):
        self.bao("Đang kiểm tra dữ liệu…")
        try:
            self.canh_bao = dc.tong_hop_canh_bao(self.ds)
        except Exception as e:
            self.bao("Lỗi kiểm tra: %s" % e)
            return
        self.tv.delete(*self.tv.get_children())
        tag = {"Nặng": "nang", "Vừa": "vua", "Nhẹ": "nhe"}
        nang = 0
        for cb in self.canh_bao[:200]:
            nang += 1 if cb["muc"] == "Nặng" else 0
            self.tv.insert("", "end", tags=(tag.get(cb["muc"], ""),),
                           values=(cb["muc"], cb["loai"], cb["hoa_don"],
                                   cb["mo_ta"]))
        if not self.canh_bao:
            self.tv.insert("", "end",
                           values=("", "", "", "Không có cảnh báo — dữ liệu sạch."))
        self.lb_cb.configure(text="%d cảnh báo (nặng: %d)"
                                  % (len(self.canh_bao), nang),
                             text_color=DO if nang else LUC)
        self.bao("Đã kiểm tra %d hóa đơn." % len(self.ds))


# ================================================== DOI CHIEU & KIEM TRA
class ManHinhDoiChieu(ManHinhNen):
    def __init__(self, master, ung_dung=None):
        super().__init__(master, ung_dung, "Đối chiếu & Kiểm tra",
                         "Kiểm tra nghiệp vụ, đối chiếu bảng kê Cổng Thuế, "
                         "so sánh hóa đơn và nạp tự động")
        self.ds, self.canh_bao, self.kq_dc = [], [], None
        self.tab = ctk.CTkTabview(self, fg_color="white",
                                  segmented_button_selected_color=XANH,
                                  segmented_button_selected_hover_color=XANH_DAM)
        self.tab.pack(fill="both", expand=True, padx=14, pady=(6, 4))
        for t in ("Kiểm tra dữ liệu", "Đối chiếu bảng kê Cổng Thuế",
                  "So sánh hai hóa đơn", "Nạp tự động"):
            self.tab.add(t)
        self._tab1(self.tab.tab("Kiểm tra dữ liệu"))
        self._tab2(self.tab.tab("Đối chiếu bảng kê Cổng Thuế"))
        self._tab3(self.tab.tab("So sánh hai hóa đơn"))
        self._tab4(self.tab.tab("Nạp tự động"))
        self.after(200, self.nap_kho)

    # ------------------------------------------------- tab 1
    def _tab1(self, cha):
        t = ctk.CTkFrame(cha, fg_color="transparent")
        t.pack(fill="x", padx=6, pady=8)
        nut(t, "Chạy kiểm tra", self.chay_kiem_tra, LUC).pack(side="left")
        nut(t, "Nạp lại kho", self.nap_kho, XAM, 120).pack(side="left", padx=8)
        nut(t, "Xuất cảnh báo ra Excel", self.xuat_canh_bao, CAM,
            200).pack(side="left")
        self.lb_cb = ctk.CTkLabel(t, text="",
                                  font=ctk.CTkFont("Segoe UI", 12, "bold"))
        self.lb_cb.pack(side="right", padx=6)
        kh, self.tv_cb = tao_bang(cha, [("muc", "Mức", 65),
                                        ("loai", "Loại", 140),
                                        ("hd", "Hóa đơn", 240),
                                        ("mo_ta", "Nội dung cần xử lý", 660)], 18)
        kh.pack(fill="both", expand=True, padx=6, pady=(0, 8))

    def chay_kiem_tra(self):
        self.bao("Đang kiểm tra…")
        try:
            self.canh_bao = dc.tong_hop_canh_bao(self.ds)
        except Exception as e:
            messagebox.showerror("Lỗi", str(e))
            return
        self.tv_cb.delete(*self.tv_cb.get_children())
        tag = {"Nặng": "nang", "Vừa": "vua", "Nhẹ": "nhe"}
        dem = {"Nặng": 0, "Vừa": 0, "Nhẹ": 0}
        for c in self.canh_bao:
            dem[c["muc"]] = dem.get(c["muc"], 0) + 1
            self.tv_cb.insert("", "end", tags=(tag.get(c["muc"], ""),),
                              values=(c["muc"], c["loai"], c["hoa_don"],
                                      c["mo_ta"]))
        if not self.canh_bao:
            self.tv_cb.insert("", "end",
                              values=("", "", "", "Không có cảnh báo nào."))
        self.lb_cb.configure(text="Nặng: %d   Vừa: %d   Nhẹ: %d"
                                  % (dem["Nặng"], dem["Vừa"], dem["Nhẹ"]),
                             text_color=DO if dem["Nặng"] else XANH_DAM)
        self.bao("Đã kiểm tra %d hóa đơn, %d cảnh báo."
                 % (len(self.ds), len(self.canh_bao)))

    def xuat_canh_bao(self):
        if not self.canh_bao:
            self.chay_kiem_tra()
        d = filedialog.asksaveasfilename(
            title="Lưu danh sách cảnh báo", defaultextension=".xlsx",
            initialfile="CanhBao_%s.xlsx"
                        % datetime.datetime.now().strftime("%Y%m%d_%H%M"),
            filetypes=[("Excel", "*.xlsx")])
        if not d:
            return
        try:
            bc.xuat_bao_cao(self.ds, d, self.canh_bao)
            if messagebox.askyesno("Xong", "Đã lưu:\n%s\n\nMở ngay?" % d):
                os.startfile(d)
            self.bao("Đã lưu %s" % d)
        except Exception as e:
            messagebox.showerror("Lỗi", str(e))

    # ------------------------------------------------- tab 2
    def _tab2(self, cha):
        t = ctk.CTkFrame(cha, fg_color="transparent")
        t.pack(fill="x", padx=6, pady=8)
        ctk.CTkLabel(t, text="Bảng kê xuất từ Cổng Thuế:").pack(side="left")
        self.o_excel = ctk.CTkEntry(t, width=460, height=32,
                                    placeholder_text="Tệp .xlsx hoặc .csv")
        self.o_excel.pack(side="left", padx=8)
        nut(t, "Chọn tệp…", self.chon_excel, XAM, 110).pack(side="left")
        nut(t, "ĐỐI CHIẾU", self.chay_doi_chieu, LUC, 130).pack(side="left",
                                                                padx=8)
        nut(t, "Xuất kết quả", self.xuat_dc, CAM, 130).pack(side="left")
        self.lb_dc = ctk.CTkLabel(cha, text="Chưa đối chiếu.",
                                  text_color=XANH_DAM,
                                  font=ctk.CTkFont("Segoe UI", 12, "bold"))
        self.lb_dc.pack(anchor="w", padx=10)
        kh, self.tv_dc = tao_bang(cha, [("nhom", "Nhóm", 180),
                                        ("ngay", "Ngày lập", 90),
                                        ("kh", "Ký hiệu", 90),
                                        ("so", "Số HĐ", 90),
                                        ("mst", "MST bán", 110),
                                        ("ten", "Tên người bán", 300),
                                        ("tien", "Tổng tiền", 110),
                                        ("ghi", "Ghi chú", 260)], 16)
        kh.pack(fill="both", expand=True, padx=6, pady=(4, 8))

    def chon_excel(self):
        d = filedialog.askopenfilename(
            title="Chọn bảng kê xuất từ Cổng Thuế",
            filetypes=[("Excel / CSV", "*.xlsx *.xls *.csv"), ("Tất cả", "*.*")])
        if d:
            self.o_excel.delete(0, "end")
            self.o_excel.insert(0, d)

    def chay_doi_chieu(self):
        d = self.o_excel.get().strip().strip('"')
        if not d or not os.path.exists(d):
            messagebox.showwarning("Thiếu tệp", "Hãy chọn tệp bảng kê trước.")
            return
        try:
            ds_cong, tt = dc.doc_bang_ke(d)
        except Exception as e:
            messagebox.showerror("Không đọc được bảng kê", str(e))
            return
        kq = dc.doi_chieu_bang_ke(self.ds, ds_cong)
        self.kq_dc = kq
        self.tv_dc.delete(*self.tv_dc.get_children())
        for h in kq["chi_co_cong"]:
            self.tv_dc.insert("", "end", tags=("nang",), values=(
                "Cổng có — máy CHƯA có", dc.ngay_vn(h["ngay_lap"]),
                h["ky_hieu_hoa_don"], h["so_hoa_don"], h["mst_nguoi_ban"],
                h["ten_nguoi_ban"], dc.fmt_tien(h["tong_tien"]),
                "Cần tải XML hóa đơn này về máy"))
        for h in kq["chi_co_may"]:
            self.tv_dc.insert("", "end", tags=("vua",), values=(
                "Máy có — cổng KHÔNG có", dc.ngay_vn(h["ngay_lap"]),
                h["ky_hieu_hoa_don"], h["so_hoa_don"], h["mst_nguoi_ban"],
                h["ten_nguoi_ban"], dc.fmt_tien(h["tong_tien"]),
                "Kiểm tra kỳ lọc trên cổng hoặc hóa đơn đã hủy"))
        for l in kq["lech_tien"]:
            a, b = l["may"], l["cong"]
            self.tv_dc.insert("", "end", tags=("vua",), values=(
                "Lệch tổng tiền", dc.ngay_vn(a["ngay_lap"]),
                a["ky_hieu_hoa_don"], a["so_hoa_don"], a["mst_nguoi_ban"],
                a["ten_nguoi_ban"], dc.fmt_tien(a["tong_tien"]),
                "Cổng ghi %s đ" % dc.fmt_tien(b["tong_tien"])))
        if not self.tv_dc.get_children():
            self.tv_dc.insert("", "end", values=("Khớp hoàn toàn", "", "", "",
                                                 "", "", "", ""))
        self.lb_dc.configure(
            text="Bảng kê %d dòng (tiêu đề dòng %d) • Khớp %d • Cổng có máy "
                 "chưa %d • Máy có cổng chưa %d • Lệch tiền %d"
                 % (tt["so_dong"], tt["dong_tieu_de"], kq["khop"],
                    len(kq["chi_co_cong"]), len(kq["chi_co_may"]),
                    len(kq["lech_tien"])))
        self.bao("Đối chiếu xong.")

    def xuat_dc(self):
        if not self.kq_dc:
            messagebox.showinfo("Chưa có dữ liệu", "Hãy chạy đối chiếu trước.")
            return
        d = filedialog.asksaveasfilename(
            title="Lưu kết quả đối chiếu", defaultextension=".xlsx",
            initialfile="DoiChieu_%s.xlsx"
                        % datetime.datetime.now().strftime("%Y%m%d_%H%M"),
            filetypes=[("Excel", "*.xlsx")])
        if not d:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            wb = Workbook()
            ws = wb.active
            ws.title = "Đối chiếu"
            ws.append(["Nhóm", "Ngày lập", "Ký hiệu", "Số HĐ", "MST bán",
                       "Tên người bán", "Tổng tiền", "Ghi chú"])
            for i, w in enumerate([26, 12, 11, 12, 15, 40, 15, 40], 1):
                ws.column_dimensions[chr(64 + i)].width = w
                o = ws.cell(row=1, column=i)
                o.font = Font(bold=True, color="FFFFFF")
                o.fill = PatternFill("solid", fgColor="1565C0")
            for iid in self.tv_dc.get_children():
                ws.append(list(self.tv_dc.item(iid)["values"]))
            ws.freeze_panes = "A2"
            wb.save(d)
            if messagebox.askyesno("Xong", "Đã lưu:\n%s\n\nMở ngay?" % d):
                os.startfile(d)
        except Exception as e:
            messagebox.showerror("Lỗi", str(e))

    # ------------------------------------------------- tab 3
    def _tab3(self, cha):
        t = ctk.CTkFrame(cha, fg_color="transparent")
        t.pack(fill="x", padx=6, pady=8)
        ctk.CTkLabel(t, text="Lọc nhanh:").pack(side="left")
        self.o_loc = ctk.CTkEntry(t, width=300, height=32,
                                  placeholder_text="số hóa đơn / ký hiệu / "
                                                   "tên người bán")
        self.o_loc.pack(side="left", padx=8)
        self.o_loc.bind("<Return>", lambda e: self.loc_ss())
        nut(t, "Lọc", self.loc_ss, XAM, 90).pack(side="left")
        nut(t, "SO SÁNH HAI DÒNG ĐANG CHỌN", self.so_sanh, LUC,
            270).pack(side="left", padx=10)
        cot = [("ngay", "Ngày lập", 95), ("kh", "Ký hiệu", 95),
               ("so", "Số HĐ", 95), ("ten", "Tên người bán", 380),
               ("tien", "Tổng tiền", 120), ("tt", "Trạng thái", 110)]
        ctk.CTkLabel(cha, text="HÓA ĐƠN A", text_color=XANH_DAM,
                     font=ctk.CTkFont("Segoe UI", 12, "bold")).pack(anchor="w",
                                                                    padx=10)
        kh1, self.tv_a = tao_bang(cha, cot, 7)
        kh1.pack(fill="both", expand=True, padx=6)
        ctk.CTkLabel(cha, text="HÓA ĐƠN B", text_color=XANH_DAM,
                     font=ctk.CTkFont("Segoe UI", 12, "bold")).pack(anchor="w",
                                                                    padx=10,
                                                                    pady=(8, 0))
        kh2, self.tv_b = tao_bang(cha, cot, 7)
        kh2.pack(fill="both", expand=True, padx=6, pady=(0, 8))
        self._bd = {}

    def loc_ss(self):
        tk_ = dc.khong_dau(self.o_loc.get())
        self.tv_a.delete(*self.tv_a.get_children())
        self.tv_b.delete(*self.tv_b.get_children())
        self._bd = {}
        n = 0
        for h in self.ds:
            chuoi = dc.khong_dau("%s %s %s %s" % (h["so_hoa_don"],
                                                  h["ky_hieu_hoa_don"],
                                                  h["ten_nguoi_ban"],
                                                  h["mst_nguoi_ban"]))
            if tk_ and tk_ not in chuoi:
                continue
            gt = (dc.ngay_vn(h["ngay_lap"]),
                  h["ky_hieu_mau_so"] + h["ky_hieu_hoa_don"], h["so_hoa_don"],
                  h["ten_nguoi_ban"], dc.fmt_tien(h["tong_tien"]),
                  h["trang_thai"])
            self._bd[("a", self.tv_a.insert("", "end", values=gt))] = h
            self._bd[("b", self.tv_b.insert("", "end", values=gt))] = h
            n += 1
            if n >= 400:
                break
        self.bao("Hiển thị %d hóa đơn để so sánh." % n)

    def so_sanh(self):
        sa, sb = self.tv_a.selection(), self.tv_b.selection()
        if not sa or not sb:
            messagebox.showinfo("Chưa chọn đủ", "Hãy chọn một dòng ở mỗi bảng.")
            return
        a, b = self._bd.get(("a", sa[0])), self._bd.get(("b", sb[0]))
        t = os.path.join(THU_MUC, "ban_the_hien")
        os.makedirs(t, exist_ok=True)
        d = os.path.join(t, "SoSanh_%s_%s_%s.html"
                         % (a["so_hoa_don"], b["so_hoa_don"],
                            datetime.datetime.now().strftime("%H%M%S")))
        try:
            bc.so_sanh_hai_hoa_don(a, b, d)
            webbrowser.open("file:///" + d.replace("\\", "/"))
            self.bao("Đã mở bản so sánh.")
        except Exception as e:
            messagebox.showerror("Lỗi", str(e))

    # ------------------------------------------------- tab 4
    def _tab4(self, cha):
        ch = theo_doi.doc_cau_hinh()
        cuon = ctk.CTkScrollableFrame(cha, fg_color="white")
        cuon.pack(fill="both", expand=True, padx=6, pady=8)

        k1 = ctk.CTkFrame(cuon, fg_color="#F4F7FB", corner_radius=8)
        k1.pack(fill="x", pady=(0, 10))
        ctk.CTkLabel(k1, text="A. Theo dõi thư mục (tự nhận tệp tải từ Cổng Thuế)",
                     font=ctk.CTkFont("Segoe UI", 13, "bold"),
                     text_color=XANH_DAM).pack(anchor="w", padx=14, pady=(10, 2))
        ctk.CTkLabel(k1, text="Phần mềm chỉ đọc tệp .xml và .zip đã có trên "
                              "máy, nhận diện trùng bằng mã băm SHA-256. "
                              "Không kết nối tới máy chủ cơ quan thuế.",
                     text_color="#5A6B7B",
                     font=ctk.CTkFont("Segoe UI", 11)).pack(anchor="w", padx=14)
        kh, self.tv_tm = tao_bang(k1, [("tm", "Thư mục đang theo dõi", 760)], 4)
        kh.pack(fill="x", padx=14, pady=8)
        for t in ch.get("thu_muc", []):
            self.tv_tm.insert("", "end", values=(t,))
        h1 = ctk.CTkFrame(k1, fg_color="transparent")
        h1.pack(fill="x", padx=14, pady=(0, 8))
        nut(h1, "Thêm thư mục…", self.them_tm, XANH, 150).pack(side="left")
        nut(h1, "Bỏ thư mục", self.bo_tm, XAM, 130).pack(side="left", padx=8)
        ctk.CTkLabel(h1, text="Chu kỳ quét (giây)").pack(side="left", padx=(20, 6))
        self.o_chu_ky = ctk.CTkEntry(h1, width=70, height=30)
        self.o_chu_ky.insert(0, str(ch.get("chu_ky_giay", 30)))
        self.o_chu_ky.pack(side="left")
        ctk.CTkLabel(h1, text="Chỉ xét tệp mới trong (ngày)").pack(side="left",
                                                                   padx=(20, 6))
        self.o_ngay = ctk.CTkEntry(h1, width=70, height=30)
        self.o_ngay.insert(0, str(ch.get("so_ngay_gan_day", 30)))
        self.o_ngay.pack(side="left")
        self.cb_de_quy = ctk.CTkCheckBox(h1, text="Quét cả thư mục con")
        self.cb_de_quy.pack(side="left", padx=18)
        if ch.get("de_quy", True):
            self.cb_de_quy.select()

        h2 = ctk.CTkFrame(k1, fg_color="transparent")
        h2.pack(fill="x", padx=14, pady=(0, 12))
        self.nut_bat = nut(h2, "BẬT THEO DÕI", self.bat_tat, LUC, 170)
        self.nut_bat.pack(side="left")
        nut(h2, "Quét ngay một lần", self.quet_ngay, XANH, 170).pack(side="left",
                                                                     padx=8)
        nut(h2, "Lưu thiết lập", self.luu_cau_hinh, XAM, 140).pack(side="left")
        nut(h2, "Xóa lịch sử mã băm", self.xoa_bam, DO, 180).pack(side="left",
                                                                  padx=8)
        self.lb_dv = ctk.CTkLabel(h2, text="Đang tắt",
                                  font=ctk.CTkFont("Segoe UI", 12, "bold"),
                                  text_color=XAM)
        self.lb_dv.pack(side="right")

        k2 = ctk.CTkFrame(cuon, fg_color="#F4F7FB", corner_radius=8)
        k2.pack(fill="x", pady=(0, 10))
        ctk.CTkLabel(k2, text="B. Hộp thư điện tử (IMAP) — lấy hóa đơn người "
                              "bán gửi qua email",
                     font=ctk.CTkFont("Segoe UI", 13, "bold"),
                     text_color=XANH_DAM).pack(anchor="w", padx=14, pady=(10, 2))
        ctk.CTkLabel(k2, text="Kết nối tới hộp thư của bạn, không phải hệ "
                              "thống thuế. Gmail/Outlook yêu cầu 'mật khẩu "
                              "ứng dụng'. Mật khẩu lưu trong Windows "
                              "Credential Manager, không ghi vào tệp cấu hình.",
                     text_color="#5A6B7B", wraplength=900, justify="left",
                     font=ctk.CTkFont("Segoe UI", 11)).pack(anchor="w", padx=14)
        g = ctk.CTkFrame(k2, fg_color="transparent")
        g.pack(fill="x", padx=14, pady=10)
        im = ch.get("imap", {})
        self.o_may = self._o(g, 0, 0, "Máy chủ IMAP", im.get("may_chu", ""), 200)
        self.o_cong = self._o(g, 0, 2, "Cổng", str(im.get("cong", 993)), 80)
        self.o_tk = self._o(g, 1, 0, "Tài khoản (email)",
                            im.get("tai_khoan", ""), 260)
        self.o_mk = self._o(g, 1, 2, "Mật khẩu ứng dụng", "", 200)
        self.o_mk.configure(show="•")
        self.o_hop = self._o(g, 2, 0, "Hộp thư", im.get("hop_thu", "INBOX"), 160)
        self.o_songay = self._o(g, 2, 2, "Lấy thư trong (ngày)",
                                str(im.get("so_ngay", 30)), 80)
        self.cb_ssl = ctk.CTkCheckBox(g, text="Dùng SSL")
        self.cb_ssl.grid(row=3, column=0, sticky="w", pady=8)
        if im.get("ssl", True):
            self.cb_ssl.select()
        mk_cu = hop_thu.doc_mat_khau(im.get("may_chu", ""),
                                     im.get("tai_khoan", ""))
        if mk_cu:
            self.o_mk.insert(0, mk_cu)

        h3 = ctk.CTkFrame(k2, fg_color="transparent")
        h3.pack(fill="x", padx=14, pady=(0, 12))
        nut(h3, "Lưu mật khẩu an toàn", self.luu_mk, XANH, 190).pack(side="left")
        nut(h3, "Kiểm tra kết nối", self.kiem_tra_imap, XAM,
            160).pack(side="left", padx=8)
        nut(h3, "TẢI ĐÍNH KÈM & NẠP", self.tai_imap, LUC,
            200).pack(side="left")
        nut(h3, "Xóa mật khẩu đã lưu", self.xoa_mk, DO, 180).pack(side="left",
                                                                  padx=8)

        ctk.CTkLabel(cuon, text="Nhật ký nạp tự động",
                     font=ctk.CTkFont("Segoe UI", 13, "bold"),
                     text_color=XANH_DAM).pack(anchor="w", padx=4)
        self.txt = tk.Text(cuon, height=12, font=("Consolas", 9), bg="white",
                           relief="solid", bd=1, wrap="word")
        self.txt.pack(fill="both", expand=True, pady=(4, 8))
        self.ghi("Sẵn sàng. Thư mục mặc định: %s" % theo_doi.thu_muc_tai_ve())
        self.after(400, self.cap_nhat_dv)

    def _o(self, g, r, c, nhan, gt, w):
        ctk.CTkLabel(g, text=nhan).grid(row=r, column=c, sticky="w",
                                        padx=(0 if c == 0 else 20, 8), pady=6)
        e = ctk.CTkEntry(g, width=w, height=30)
        e.grid(row=r, column=c + 1, sticky="w", pady=6)
        if gt:
            e.insert(0, gt)
        return e

    def ghi(self, t):
        try:
            self.txt.insert("end", "%s  %s\n"
                            % (datetime.datetime.now().strftime("%H:%M:%S"), t))
            self.txt.see("end")
        except Exception:
            pass
        self.bao(t)

    def them_tm(self):
        d = filedialog.askdirectory(title="Chọn thư mục cần theo dõi")
        if d:
            self.tv_tm.insert("", "end", values=(d,))
            self.luu_cau_hinh()

    def bo_tm(self):
        for i in self.tv_tm.selection():
            self.tv_tm.delete(i)
        self.luu_cau_hinh()

    def _cau_hinh_hien_tai(self):
        ch = theo_doi.doc_cau_hinh()
        ch["thu_muc"] = [self.tv_tm.item(i)["values"][0]
                         for i in self.tv_tm.get_children()]
        try:
            ch["chu_ky_giay"] = max(5, int(self.o_chu_ky.get()))
        except ValueError:
            pass
        try:
            ch["so_ngay_gan_day"] = max(1, int(self.o_ngay.get()))
        except ValueError:
            pass
        ch["de_quy"] = bool(self.cb_de_quy.get())
        ch["imap"].update({
            "may_chu": self.o_may.get().strip(),
            "cong": int(self.o_cong.get() or 993),
            "ssl": bool(self.cb_ssl.get()),
            "tai_khoan": self.o_tk.get().strip(),
            "hop_thu": self.o_hop.get().strip() or "INBOX",
            "so_ngay": int(self.o_songay.get() or 30)})
        return ch

    def luu_cau_hinh(self):
        try:
            ch = self._cau_hinh_hien_tai()
            theo_doi.ghi_cau_hinh(ch)
            self.ghi("Đã lưu thiết lập nạp tự động.")
        except Exception as e:
            messagebox.showerror("Lỗi", str(e))

    def _dich_vu(self):
        return theo_doi.lay_dich_vu(goi_ui=lambda f: self.after(0, f),
                                    khi_xong=self._sau_khi_nap,
                                    bao=self.ghi)

    def _sau_khi_nap(self, kq):
        self.nap_kho()
        if self.ung_dung and hasattr(self.ung_dung, "man_hinh"):
            mh = self.ung_dung.man_hinh.get("trang_chu")
            if mh is not None and hasattr(mh, "lam_moi"):
                try:
                    mh.lam_moi()
                except Exception:
                    pass

    def bat_tat(self):
        self.luu_cau_hinh()
        dv = self._dich_vu()
        if dv.dang_chay():
            dv.tat()
            self.ghi("Đã tắt theo dõi thư mục.")
        else:
            dv.bat()
            self.ghi("Đã bật theo dõi thư mục, chu kỳ %s giây."
                     % self.o_chu_ky.get())
        ch = theo_doi.doc_cau_hinh()
        ch["bat_theo_doi"] = dv.dang_chay()
        theo_doi.ghi_cau_hinh(ch)
        self.cap_nhat_dv()

    def cap_nhat_dv(self):
        dv = theo_doi.DICH_VU
        chay = bool(dv and dv.dang_chay())
        self.lb_dv.configure(text="Đang chạy" if chay else "Đang tắt",
                             text_color=LUC if chay else XAM)
        self.nut_bat.configure(text="TẮT THEO DÕI" if chay else "BẬT THEO DÕI",
                               fg_color=DO if chay else LUC)
        self.after(3000, self.cap_nhat_dv)

    def quet_ngay(self):
        self.luu_cau_hinh()
        self.ghi("Đang quét thư mục…")
        dv = self._dich_vu()
        dv.quet_ngay()

    def xoa_bam(self):
        if messagebox.askyesno("Xác nhận",
                               "Xóa lịch sử mã băm? Các tệp cũ sẽ được quét "
                               "lại từ đầu (hóa đơn trùng vẫn không bị nhân "
                               "đôi nhờ khóa MST + ký hiệu + số)."):
            theo_doi.xoa_lich_su_bam()
            self.ghi("Đã xóa lịch sử mã băm.")

    def luu_mk(self):
        try:
            hop_thu.luu_mat_khau(self.o_may.get().strip(),
                                 self.o_tk.get().strip(), self.o_mk.get())
            self.luu_cau_hinh()
            self.ghi("Đã lưu mật khẩu vào Windows Credential Manager.")
        except Exception as e:
            messagebox.showerror("Lỗi", str(e))

    def xoa_mk(self):
        hop_thu.xoa_mat_khau(self.o_may.get().strip(), self.o_tk.get().strip())
        self.o_mk.delete(0, "end")
        self.ghi("Đã xóa mật khẩu đã lưu.")

    def kiem_tra_imap(self):
        ch = self._cau_hinh_hien_tai()["imap"]
        mk = self.o_mk.get() or hop_thu.doc_mat_khau(ch["may_chu"],
                                                     ch["tai_khoan"])
        self.ghi("Đang kiểm tra kết nối hộp thư…")

        def chay():
            ok, tb = hop_thu.kiem_tra_ket_noi(ch, mk)
            self.after(0, lambda: self.ghi(("OK — " if ok else "LỖI — ") + tb))
        threading.Thread(target=chay, daemon=True).start()

    def tai_imap(self):
        ch = self._cau_hinh_hien_tai()
        theo_doi.ghi_cau_hinh(ch)
        im = ch["imap"]
        mk = self.o_mk.get() or hop_thu.doc_mat_khau(im["may_chu"],
                                                     im["tai_khoan"])
        if not im["may_chu"] or not im["tai_khoan"] or not mk:
            messagebox.showwarning("Thiếu thông tin",
                                   "Hãy điền máy chủ, tài khoản và mật khẩu "
                                   "ứng dụng.")
            return
        self.ghi("Đang tải đính kèm từ hộp thư…")

        def chay():
            try:
                kq = hop_thu.tai_dinh_kem(im, mk,
                                          bao=lambda t: self.after(
                                              0, lambda: self.ghi(t)))
            except Exception as e:
                loi = str(e)
                self.after(0, lambda: self.ghi("LỖI: %s" % loi))
                return

            def nap():
                if not kq["tep"]:
                    self.ghi("Không có đính kèm mới (bỏ qua %d thư)."
                             % kq["bo_qua"])
                    return
                r = theo_doi.nap_nhieu_tep(kq["tep"], bao=self.ghi)
                self.ghi("Hộp thư: %d thư có đính kèm, nạp %d tệp, thêm %d "
                         "hóa đơn, trùng %d, lỗi %d."
                         % (kq["thu"], r["tep"], r["them"], r["trung"],
                            r["loi"]))
                self._sau_khi_nap(r)
            self.after(0, nap)
        threading.Thread(target=chay, daemon=True).start()

    # ------------------------------------------------- chung
    def nap_kho(self):
        try:
            self.ds = dc.doc_hoa_don()
        except Exception as e:
            self.bao("Không đọc được kho: %s" % e)
            self.ds = []
        self.bao("Đã nạp %d hóa đơn." % len(self.ds))
        try:
            self.loc_ss()
        except Exception:
            pass


# ==================================================== BAO CAO & KE KHAI
class ManHinhBaoCao(ManHinhNen):
    def __init__(self, master, ung_dung=None):
        super().__init__(master, ung_dung, "Báo cáo & Kê khai",
                         "Chi tiết từng mặt hàng, bảng kê mua vào – bán ra, "
                         "thống kê và gói lưu trữ theo kỳ")
        self.ds, self.dong = [], []
        self.tab = ctk.CTkTabview(self, fg_color="white",
                                  segmented_button_selected_color=XANH,
                                  segmented_button_selected_hover_color=XANH_DAM)
        self.tab.pack(fill="both", expand=True, padx=14, pady=(6, 4))
        self.tab.add("Chi tiết mặt hàng & Kê khai GTGT")
        self.tab.add("Bảng kê tổng hợp & Gói lưu trữ")
        self._tab1(self.tab.tab("Chi tiết mặt hàng & Kê khai GTGT"))
        self._tab2(self.tab.tab("Bảng kê tổng hợp & Gói lưu trữ"))
        self.after(200, self.nap_kho)

    def _tham_so(self, cha):
        k = ctk.CTkFrame(cha, fg_color="#F4F7FB", corner_radius=8)
        k.pack(fill="x", padx=6, pady=8)
        ctk.CTkLabel(k, text="MST đơn vị kê khai").grid(row=0, column=0,
                                                         padx=12, pady=10,
                                                         sticky="w")
        o_mst = ctk.CTkEntry(k, width=170, height=30,
                             placeholder_text="VD: 0101234567")
        o_mst.grid(row=0, column=1, sticky="w")
        ctk.CTkLabel(k, text="Từ ngày").grid(row=0, column=2, padx=(22, 6),
                                             sticky="w")
        o_tu = ctk.CTkEntry(k, width=120, height=30,
                            placeholder_text="dd/mm/yyyy")
        o_tu.grid(row=0, column=3, sticky="w")
        ctk.CTkLabel(k, text="Đến ngày").grid(row=0, column=4, padx=(22, 6),
                                              sticky="w")
        o_den = ctk.CTkEntry(k, width=120, height=30,
                             placeholder_text="dd/mm/yyyy")
        o_den.grid(row=0, column=5, sticky="w")
        ctk.CTkLabel(k, text="Để trống hai ô ngày nếu lấy toàn bộ kho. MST "
                             "dùng để tách Mua vào / Bán ra.",
                     text_color="#5A6B7B",
                     font=ctk.CTkFont("Segoe UI", 11)).grid(
            row=1, column=0, columnspan=6, sticky="w", padx=12, pady=(0, 10))
        try:
            import json
            with open(os.path.join(THU_MUC, "config.json"), "r",
                      encoding="utf-8") as f:
                c = json.load(f)
            for kk, vv in c.items():
                if "mst" in kk.lower() and str(vv).strip():
                    o_mst.insert(0, str(vv).strip())
                    break
        except Exception:
            pass
        return o_mst, o_tu, o_den

    def _tab1(self, cha):
        self.o_mst, self.o_tu, self.o_den = self._tham_so(cha)
        t = ctk.CTkFrame(cha, fg_color="transparent")
        t.pack(fill="x", padx=6)
        nut(t, "Xem trước chi tiết", self.xem_truoc, XAM, 170).pack(side="left")
        nut(t, "XUẤT EXCEL ĐẦY ĐỦ (9 sheet)", self.xuat_excel, LUC,
            250).pack(side="left", padx=8)
        nut(t, "Xuất CSV bảng phẳng", self.xuat_csv, CAM,
            180).pack(side="left")
        kh, self.tv = tao_bang(cha, [("ngay", "Ngày", 85),
                                     ("loai", "Loại GD", 80),
                                     ("so", "Số HĐ", 80),
                                     ("dt", "Đối tác", 230),
                                     ("hang", "Tên hàng hóa, dịch vụ", 320),
                                     ("dvt", "ĐVT", 60),
                                     ("sl", "Số lượng", 80),
                                     ("dg", "Đơn giá", 100),
                                     ("tt", "Thành tiền", 110),
                                     ("ts", "Thuế suất", 150),
                                     ("thue", "Tiền thuế", 100),
                                     ("ng", "Nguồn dòng", 130)], 15)
        kh.pack(fill="both", expand=True, padx=6, pady=8)

    def _tab2(self, cha):
        self.o_mst2, self.o_tu2, self.o_den2 = self._tham_so(cha)
        t = ctk.CTkFrame(cha, fg_color="transparent")
        t.pack(fill="x", padx=6)
        nut(t, "Xem trước số liệu", self.xem_truoc2, XAM, 170).pack(side="left")
        nut(t, "XUẤT BẢNG KÊ EXCEL", self.xuat_bang_ke, LUC,
            210).pack(side="left", padx=8)
        nut(t, "XUẤT GÓI LƯU TRỮ", self.xuat_goi, CAM, 190).pack(side="left")
        self.txt2 = tk.Text(cha, height=20, font=("Consolas", 10), bg="white",
                            relief="solid", bd=1, wrap="word")
        self.txt2.pack(fill="both", expand=True, padx=6, pady=8)
        self.txt2.insert("1.0",
                         "Bảng kê Excel gồm 6 sheet: Thông tin, Danh sách, "
                         "Dòng hàng, Theo nhà cung cấp, Theo thuế suất, "
                         "Cảnh báo.\nGói lưu trữ tạo thư mục chứa XML gốc, "
                         "bản thể hiện HTML và danh mục Excel.")

    # --------------------------------------------------- xu ly
    def nap_kho(self):
        try:
            self.ds = dc.doc_hoa_don()
        except Exception as e:
            self.bao("Không đọc được kho: %s" % e)
            self.ds = []
        self.bao("Đã nạp %d hóa đơn." % len(self.ds))

    def _ky(self, o_tu, o_den, o_mst):
        tu = dc.ngay_chuan(o_tu.get().strip())
        den = dc.ngay_chuan(o_den.get().strip())
        ds = [h for h in self.ds
              if (not tu or (h["ngay_lap"] and h["ngay_lap"] >= tu))
              and (not den or (h["ngay_lap"] and h["ngay_lap"] <= den))]
        return ds, tu, den, o_mst.get().strip()

    def xem_truoc(self):
        ds, tu, den, mst = self._ky(self.o_tu, self.o_den, self.o_mst)
        if not ds:
            messagebox.showinfo("Không có dữ liệu", "Kỳ đã chọn không có hóa đơn.")
            return
        self.bao("Đang dựng bảng chi tiết…")
        self.dong = bct.bang_phang(ds, mst)
        self.tv.delete(*self.tv.get_children())
        for r in self.dong[:1500]:
            tag = ("vua",) if r["nguon_dong"].startswith("Suy") else (
                ("ban",) if r["loai_gd"] == "Bán ra" else ())
            self.tv.insert("", "end", tags=tag, values=(
                r["ngay_lap"], r["loai_gd"], r["so_hd"],
                r["ten_mua"] if r["loai_gd"] == "Bán ra" else r["ten_ban"],
                r["ten_hang"], r["dvt"], "%g" % r["so_luong"],
                dc.fmt_tien(r["don_gia"]), dc.fmt_tien(r["thanh_tien"]),
                r["ts_nhan"], dc.fmt_tien(r["tien_thue_dong"]),
                r["nguon_dong"]))
        mua = len([r for r in self.dong if r["loai_gd"] == "Mua vào"])
        ban = len([r for r in self.dong if r["loai_gd"] == "Bán ra"])
        self.bao("%d hóa đơn → %d dòng mặt hàng (mua vào %d, bán ra %d)."
                 % (len(ds), len(self.dong), mua, ban))

    def xuat_excel(self):
        ds, tu, den, mst = self._ky(self.o_tu, self.o_den, self.o_mst)
        if not ds:
            messagebox.showinfo("Không có dữ liệu", "Kỳ đã chọn không có hóa đơn.")
            return
        d = filedialog.asksaveasfilename(
            title="Lưu báo cáo chi tiết", defaultextension=".xlsx",
            initialfile="BaoCaoChiTiet_%s.xlsx"
                        % datetime.datetime.now().strftime("%Y%m%d_%H%M"),
            filetypes=[("Excel", "*.xlsx")])
        if not d:
            return
        self.bao("Đang lập báo cáo…")

        def chay():
            try:
                kq = bct.xuat_chi_tiet_day_du(ds, d, mst, tu, den)
            except Exception as e:
                v = "%s\n\n%s" % (e, traceback.format_exc()[-800:])
                self.after(0, lambda: messagebox.showerror("Lỗi", v))
                return

            def xong():
                self.bao("Đã lưu %s" % kq["tep"])
                if messagebox.askyesno(
                        "Hoàn tất",
                        "Đã lập báo cáo:\n%s\n\nHóa đơn %d — dòng mặt hàng %d "
                        "(mua vào %d, bán ra %d).\n\nMở ngay?"
                        % (kq["tep"], kq["so_hoa_don"], kq["so_dong"],
                           kq["mua"], kq["ban"])):
                    os.startfile(kq["tep"])
            self.after(0, xong)
        threading.Thread(target=chay, daemon=True).start()

    def xuat_csv(self):
        ds, tu, den, mst = self._ky(self.o_tu, self.o_den, self.o_mst)
        if not ds:
            messagebox.showinfo("Không có dữ liệu", "Kỳ đã chọn không có hóa đơn.")
            return
        d = filedialog.asksaveasfilename(
            title="Lưu bảng phẳng CSV", defaultextension=".csv",
            initialfile="ChiTietMatHang_%s.csv"
                        % datetime.datetime.now().strftime("%Y%m%d_%H%M"),
            filetypes=[("CSV (dấu chấm phẩy)", "*.csv")])
        if not d:
            return
        try:
            bct.xuat_csv(bct.bang_phang(ds, mst), d)
            if messagebox.askyesno("Xong", "Đã lưu:\n%s\n\nMở ngay?" % d):
                os.startfile(d)
        except Exception as e:
            messagebox.showerror("Lỗi", str(e))

    def xem_truoc2(self):
        ds, tu, den, _ = self._ky(self.o_tu2, self.o_den2, self.o_mst2)
        cb = dc.tong_hop_canh_bao(ds)
        tong = {}
        for h in ds:
            g = tong.setdefault(h["ten_nguoi_ban"] or h["mst_nguoi_ban"], [0, 0])
            g[0] += 1
            g[1] += h["tong_tien"]
        self.txt2.delete("1.0", "end")
        self.txt2.insert("1.0",
                         "KỲ: %s  →  %s\nSố hóa đơn      : %d\n"
                         "Số nhà cung cấp : %d\nTổng tiền hàng  : %s đ\n"
                         "Tổng thuế GTGT  : %s đ\nTổng thanh toán : %s đ\n"
                         "Số cảnh báo     : %d (nặng %d)\n\n"
                         "Nhà cung cấp lớn nhất kỳ này:\n"
                         % (dc.ngay_vn(tu) or "đầu kỳ",
                            dc.ngay_vn(den) or "cuối kỳ", len(ds), len(tong),
                            dc.fmt_tien(sum(h["tien_hang"] for h in ds)),
                            dc.fmt_tien(sum(h["thue_gtgt"] for h in ds)),
                            dc.fmt_tien(sum(h["tong_tien"] for h in ds)),
                            len(cb),
                            len([c for c in cb if c["muc"] == "Nặng"])))
        for ten, g in sorted(tong.items(), key=lambda x: -x[1][1])[:15]:
            self.txt2.insert("end", "  %-46s %3d HĐ %16s đ\n"
                             % (ten[:46], g[0], dc.fmt_tien(g[1])))
        self.bao("Xem trước %d hóa đơn." % len(ds))

    def xuat_bang_ke(self):
        ds, tu, den, _ = self._ky(self.o_tu2, self.o_den2, self.o_mst2)
        if not ds:
            messagebox.showinfo("Không có dữ liệu", "Kỳ đã chọn không có hóa đơn.")
            return
        d = filedialog.asksaveasfilename(
            title="Lưu bảng kê", defaultextension=".xlsx",
            initialfile="BangKeHoaDon_%s.xlsx"
                        % datetime.datetime.now().strftime("%Y%m%d_%H%M"),
            filetypes=[("Excel", "*.xlsx")])
        if not d:
            return
        try:
            bc.xuat_bao_cao(ds, d, dc.tong_hop_canh_bao(ds), tu, den)
            if messagebox.askyesno("Xong", "Đã lưu:\n%s\n\nMở ngay?" % d):
                os.startfile(d)
            self.bao("Đã lưu bảng kê.")
        except Exception as e:
            messagebox.showerror("Lỗi", str(e))

    def xuat_goi(self):
        ds, tu, den, _ = self._ky(self.o_tu2, self.o_den2, self.o_mst2)
        if not ds:
            messagebox.showinfo("Không có dữ liệu", "Kỳ đã chọn không có hóa đơn.")
            return
        t = filedialog.askdirectory(title="Chọn nơi đặt gói lưu trữ")
        if not t:
            return
        self.bao("Đang tạo gói lưu trữ…")

        def chay():
            try:
                kq = bc.xuat_goi_luu_tru(
                    ds, t, tu, den,
                    bao_tien_do=lambda i, n: self.after(
                        0, lambda: self.bao("Đang xử lý %d/%d…" % (i, n))))
            except Exception as e:
                v = str(e)
                self.after(0, lambda: messagebox.showerror("Lỗi", v))
                return

            def xong():
                self.bao("Đã tạo gói: %s" % kq["thu_muc"])
                if messagebox.askyesno(
                        "Hoàn tất",
                        "Đã tạo gói lưu trữ:\n%s\n\nHóa đơn %d — XML gốc %d — "
                        "thiếu XML %d — HTML %d\n\nMở thư mục?"
                        % (kq["thu_muc"], len(ds), kq["xml"], kq["thieu_xml"],
                           kq["html"])):
                    os.startfile(kq["thu_muc"])
            self.after(0, xong)
        threading.Thread(target=chay, daemon=True).start()


LOP_MAN_HINH = {"trang_chu": TrangChuMoi, "doi_chieu": ManHinhDoiChieu,
                "bao_cao": ManHinhBaoCao}
